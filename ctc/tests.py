import io

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from accounts.models import AppKey, PortalSession, PortalUser

from .calculator import calculate, state_for, vpp_percent
from .columns import COLUMNS, OPTION_COLUMNS
from .models import CtcCandidate, CtcSettings, CtcUpload

BASE = '/api/ctc'
HEADINGS = [h for h, *_ in COLUMNS]
OPTION_HEADINGS = [h for h, *_ in OPTION_COLUMNS]


def _xlsx(rows, headings=HEADINGS, name='candidates.xlsx'):
    wb = Workbook()
    ws = wb.active
    ws.append(headings)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(name, buf.getvalue())


# Candidate Name, Category, Department, Designation, Grade, Cadre, Work Location, CTC
FIELDS = [f for _h, f, *_ in COLUMNS]
BASE_ROW = dict(candidate_name='Rahul Sharma', category='Sales',
                department='General Trade', designation='Area Sales Manager', grade='M1', cadre='M',
                number_of_children=2, work_location='Delhi HO', ctc=1200000,
                tax_regime='New Regime')


def _row(**over):
    """A template row by column name, in whatever order COLUMNS has."""
    d = {**BASE_ROW, **over}
    return [d[f] for f in FIELDS]


ROW = _row()


class CtcUploads(TestCase):
    def setUp(self):
        self.admin = PortalUser.objects.create(email='a@apisindia.com', employee_code='A1',
                                               name='Admin', is_superadmin=True)
        # Alice and Bob are P & C: the CTC Calculator is granted to them from
        # Admin Console. Being signed in is not enough on its own.
        self.alice = PortalUser.objects.create(email='alice@apisindia.com', employee_code='E1',
                                               name='Alice', app_access=[AppKey.CTC])
        self.bob = PortalUser.objects.create(email='bob@apisindia.com', employee_code='E2',
                                             name='Bob', app_access=[AppKey.CTC])
        # Carol is an ordinary employee with no tools granted.
        self.carol = PortalUser.objects.create(email='carol@apisindia.com', employee_code='E3',
                                               name='Carol')

    def _auth(self, user):
        return {'HTTP_AUTHORIZATION': f'Bearer {PortalSession.start(user)}'}

    def _upload(self, user, f):
        return self.client.post(f'{BASE}/uploads/', {'file': f}, **self._auth(user))

    def test_template_has_exactly_the_columns_the_uploader_reads(self):
        r = self.client.get(f'{BASE}/template/', **self._auth(self.alice))
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(io.BytesIO(r.content)).worksheets[0]
        self.assertEqual([c.value for c in ws[1]], HEADINGS + OPTION_HEADINGS)
        self.assertEqual(ws.max_row, 1, 'the upload sheet must hold nothing but headings')
        lists = [dv.formula1 for dv in ws.data_validations.dataValidation]
        self.assertIn('"Sales,HO,Plant"', lists)
        self.assertIn('"D,C,M,O,W"', lists)
        self.assertIn('"0,1,2"', lists)                                     # Number of Children
        self.assertIn('IF(N(INDEX($I:$I,ROW()))<=1200000,Lists!$C$1,Lists!$D$1:$D$3)', lists)   # Tax Regime
        # Gross Earnings: Yes/No, but an empty list (and greyed out) for O and W.
        self.assertIn('IF(OR(INDEX($F:$F,ROW())="O",INDEX($F:$F,ROW())="W",N(INDEX($I:$I,ROW()))<=1200000),'
                      'Lists!$C$1,Lists!$B$1:$B$2)', lists)
        self.assertTrue(ws.conditional_formatting)

    def test_gross_earnings_choices_are_stored_and_ignored_for_o_w_and_up_to_12_lakh(self):
        heads = HEADINGS + OPTION_HEADINGS
        # lta, meal, telephone, fuel, driver, health, books, uniform, car
        m = _row(ctc=1500000) + ['Yes', 'no', '', 'N', 'yes', 'No', '', 'Yes', 'Yes']
        o = _row(candidate_name='Ravi', grade='O3', cadre='O', ctc=1500000) + ['Yes'] * 9
        small = _row(candidate_name='Sita', ctc=1200000) + ['Yes'] * 9
        r = self._upload(self.alice, _xlsx([m, o, small], headings=heads))
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(CtcCandidate.objects.get(candidate_name='Rahul Sharma').gross_earnings,
                         {'lta': True, 'meal': False, 'fuel': False, 'driver': True,
                          'health': False, 'uniform': True, 'car': True})
        self.assertEqual(CtcCandidate.objects.get(candidate_name='Ravi').gross_earnings, {})
        self.assertEqual(CtcCandidate.objects.get(candidate_name='Sita').gross_earnings, {})
        bad = _row(ctc=1500000) + ['Yes', 'maybe']
        r = self._upload(self.alice, _xlsx([bad], headings=heads))
        self.assertEqual(r.status_code, 400)
        self.assertIn('Meal Voucher Reimbursement "maybe" should be Yes or No', str(r.json()))

    def test_everything_needs_sign_in(self):
        for url in ('template/', 'uploads/', 'settings/', 'candidates/1/breakdown/', 'candidates/1/pdf/'):
            self.assertEqual(self.client.get(f'{BASE}/{url}').status_code, 401, url)
        self.assertEqual(self.client.post(f'{BASE}/uploads/', {'file': _xlsx([ROW])}).status_code, 401)

    def test_upload_stores_every_row(self):
        second = _row(candidate_name='Priya Verma', category='ho', department='Finance',
                      designation='Manager', grade='m3', cadre='m',
                      work_location='Roorkee', ctc='12,50,000')
        r = self._upload(self.alice, _xlsx([ROW, [], second]))
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(r.json()['rowCount'], 2)
        rahul = CtcCandidate.objects.get(candidate_name='Rahul Sharma')
        self.assertEqual((rahul.employee_code, rahul.date_of_joining), ('', None))   # not on the template
        self.assertEqual(rahul.number_of_children, 2)
        priya = CtcCandidate.objects.get(candidate_name='Priya Verma')
        self.assertEqual((priya.category, priya.grade, priya.cadre), ('HO', 'M3', 'M'))
        self.assertEqual(str(priya.ctc), '1250000.00')
        self.assertEqual(priya.row_number, 4)

    def test_sales_rows_need_a_sales_department(self):
        rows = [_row(department='general trade'), _row(candidate_name='B', department='Alternate Channels'),
                _row(candidate_name='C', department='BOTH'),
                _row(candidate_name='D', category='HO', department='Finance', designation='Manager')]
        self.assertEqual(self._upload(self.alice, _xlsx(rows)).status_code, 201)
        self.assertEqual([c.department for c in CtcCandidate.objects.order_by('row_number')],
                         ['General Trade', 'Alternate Channels', 'Both', 'Finance'])
        r = self._upload(self.alice, _xlsx([_row(department='ASM')]))
        self.assertEqual(r.status_code, 400)
        self.assertIn('for Sales pick GENERAL TRADE, ALTERNATE CHANNELS or BOTH', str(r.json()))

    def test_tax_regime_is_read_and_optional(self):
        rows = [_row(tax_regime='old', ctc=1500000),
                _row(candidate_name='B', tax_regime='BOTH REGIME', ctc=1500000),
                _row(candidate_name='C', tax_regime='', ctc=1500000),
                _row(candidate_name='D', tax_regime='New Regime', ctc=1200000)]   # up to 12 lakh: ignored
        self.assertEqual(self._upload(self.alice, _xlsx(rows)).status_code, 201)
        self.assertEqual([c.tax_regime for c in CtcCandidate.objects.order_by('row_number')],
                         ['Old Regime', 'Both Regime', '', ''])
        # 'N/A' (from an older template) is read as blank, not an error.
        self.assertEqual(self._upload(self.alice, _xlsx([_row(tax_regime='N/A', ctc=1500000)])).status_code, 201)
        r = self._upload(self.alice, _xlsx([_row(tax_regime='flat tax')]))
        self.assertIn('Tax Regime "flat tax"', str(r.json()))
        # An older template without the column still uploads.
        old = [h for h in HEADINGS if h != 'Tax Regime']
        self.assertEqual(self._upload(self.alice, _xlsx([ROW[:-1]], headings=old)).status_code, 201)

    def test_an_old_template_with_employee_code_and_date_of_joining_still_uploads(self):
        heads = ['Employee Code', *HEADINGS[:6], 'Date of Joining', *HEADINGS[6:]]
        row = ['SL00225', *ROW[:6], '01-07-2026', *ROW[6:]]
        self.assertEqual(self._upload(self.alice, _xlsx([row], headings=heads)).status_code, 201)

    def test_headings_are_matched_loosely(self):
        loose = ['candidate name ', 'CATEGORY', 'department', 'designation', 'Grade',
                 'cadre', 'number of children', 'work_location', 'ctc', 'TAX_REGIME']
        self.assertEqual(self._upload(self.alice, _xlsx([ROW], headings=loose)).status_code, 201)

    def test_a_missing_column_is_named(self):
        keep = HEADINGS[:5]
        r = self._upload(self.alice, _xlsx([ROW[:5]], headings=keep))
        self.assertEqual(r.status_code, 400)
        for heading in ('Cadre', 'Work Location', 'CTC'):
            self.assertIn(heading, r.json()['error'])

    def test_any_bad_row_and_nothing_is_stored(self):
        bad = [
            _row(candidate_name=''),                                        # row 3: no name
            _row(candidate_name='Ravi', ctc='twelve lakh'),                   # row 4: CTC text
            _row(candidate_name='Amit', ctc=None),                            # row 5: no CTC
            _row(candidate_name='Neha', category='Marketing'),                # row 6: category
            _row(candidate_name='Om', category='HO', grade='O4', cadre='Officer'),  # row 7: reads as O
            _row(candidate_name='Raj', category='HO', grade='O4', cadre='M'),  # row 8: grade/cadre clash
            _row(candidate_name='Sia', category='HO', grade='', cadre=''),    # row 9: no grade/cadre
            _row(candidate_name='Uma', number_of_children='two'),            # row 10: bad count
            _row(candidate_name='Ved', number_of_children=3),                # row 11: more than two
        ]
        r = self._upload(self.alice, _xlsx([ROW, *bad]))
        self.assertEqual(r.status_code, 400)
        p = ' | '.join(r.json()['problems'])
        for expected in ('Row 3: Candidate Name is empty', 'Row 4: CTC "twelve lakh" is not a number',
                         'Row 5: CTC is empty', 'Row 6: Category "Marketing"',
                         'Row 8: Grade O4 does not belong to cadre M',
                         'Row 9: Grade is empty', 'Row 9: Cadre is empty',
                         'Row 10: Number of Children "two"',
                         'Row 11: Number of Children "3" should be 0, 1 or 2'):
            self.assertIn(expected, p)
        self.assertNotIn('Row 7', p, '"Officer" should read as cadre O')
        self.assertFalse(CtcUpload.objects.exists())

    def test_rejects_files_that_are_not_xlsx(self):
        csv = SimpleUploadedFile('c.csv', b'Candidate Name\nRahul\n')
        self.assertEqual(self._upload(self.alice, csv).status_code, 400)
        fake = SimpleUploadedFile('c.xlsx', b'not a zip')
        self.assertEqual(self._upload(self.alice, fake).status_code, 400)

    def test_people_see_their_own_uploads_and_admin_sees_all(self):
        self._upload(self.alice, _xlsx([ROW]))
        self._upload(self.bob, _xlsx([ROW]))
        mine = self.client.get(f'{BASE}/uploads/', **self._auth(self.alice)).json()
        self.assertEqual(len(mine), 1)
        self.assertEqual(mine[0]['candidates'][0]['candidateName'], 'Rahul Sharma')
        self.assertEqual(mine[0]['candidates'][0]['category'], 'Sales')
        self.assertEqual(mine[0]['candidates'][0]['ctc'], '1200000.00')
        self.assertEqual(len(self.client.get(f'{BASE}/uploads/', **self._auth(self.admin)).json()), 2)

    def test_reset_needs_the_word_and_erases_only_what_the_caller_sees(self):
        self._upload(self.alice, _xlsx([ROW, ROW]))
        self._upload(self.bob, _xlsx([ROW]))
        url = f'{BASE}/reset/'
        self.assertEqual(self.client.post(url, {}).status_code, 401)
        self.assertEqual(self.client.post(url, {'confirm': 'yes'}, **self._auth(self.alice)).status_code, 400)
        self.assertEqual(CtcUpload.objects.count(), 2)

        r = self.client.post(url, {'confirm': 'RESET'}, **self._auth(self.alice))
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()['deleted'], {'uploads': 1, 'candidates': 2})
        self.assertEqual(list(CtcUpload.objects.values_list('uploaded_by', flat=True)), [self.bob.id],
                         "Alice's reset touched Bob's upload")

        self.client.post(url, {'confirm': 'reset'}, **self._auth(self.admin))
        self.assertFalse(CtcCandidate.objects.exists())

    def test_only_the_uploader_or_admin_removes_an_upload(self):
        self._upload(self.alice, _xlsx([ROW]))
        u = CtcUpload.objects.get()
        self.assertEqual(self.client.delete(f'{BASE}/uploads/{u.id}/', **self._auth(self.bob)).status_code, 403)
        self.assertEqual(self.client.delete(f'{BASE}/uploads/{u.id}/', **self._auth(self.alice)).status_code, 200)
        self.assertFalse(CtcCandidate.objects.exists())

    # -- breakdown, PDF, settings ---------------------------------------------
    def test_breakdown_and_pdf_for_the_uploader_only(self):
        self._upload(self.alice, _xlsx([ROW]))
        c = CtcCandidate.objects.get()
        r = self.client.get(f'{BASE}/candidates/{c.id}/breakdown/', **self._auth(self.alice))
        self.assertEqual(r.status_code, 200, r.content)
        b = r.json()['breakdown']
        self.assertEqual(b['vppPercent'], 10)
        self.assertEqual(b['totalCtcAnnual'], 1200000)

        pdf = self.client.get(f'{BASE}/candidates/{c.id}/pdf/', **self._auth(self.alice))
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf['Content-Type'], 'application/pdf')
        self.assertTrue(pdf.content.startswith(b'%PDF-'))
        self.assertIn('CTC_Breakup_Rahul_Sharma.pdf', pdf['Content-Disposition'])

        for url in ('breakdown/', 'pdf/'):
            self.assertEqual(self.client.get(f'{BASE}/candidates/{c.id}/{url}', **self._auth(self.bob)).status_code,
                             404, f"Bob could open Alice's candidate via {url}")
        self.assertEqual(self.client.get(f'{BASE}/candidates/{c.id}/pdf/', **self._auth(self.admin)).status_code, 200)

    def test_only_a_superadmin_edits_settings_and_they_feed_the_breakdown(self):
        url = f'{BASE}/settings/'
        self.assertFalse(self.client.get(url, **self._auth(self.alice)).json()['canEdit'])
        r = self.client.put(url, {'gratuityPercent': 4}, content_type='application/json',
                            **self._auth(self.alice))
        self.assertEqual(r.status_code, 403)
        r = self.client.put(url, {'gratuityPercent': '4.81'},
                            content_type='application/json', **self._auth(self.admin))
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(str(CtcSettings.current().gratuity_percent), '4.81')
        bad = self.client.put(url, {'gratuityPercent': 'lots'}, content_type='application/json',
                              **self._auth(self.admin))
        self.assertEqual(bad.status_code, 400)

        self._upload(self.alice, _xlsx([ROW]))
        b = self.client.get(f'{BASE}/candidates/{CtcCandidate.objects.get().id}/breakdown/',
                            **self._auth(self.alice)).json()['breakdown']
        grat = next(x for x in b['benefits'] if x['label'] == 'Gratuity')
        basic = b['earnings'][0]['monthly']
        self.assertEqual(grat['monthly'], round(basic * 0.0481))


class Calculator(TestCase):
    def setUp(self):
        self.s = CtcSettings.current()

    def calc(self, **kw):
        base = dict(ctc_annual=1200000, cadre='M', grade='M1', category='Sales',
                    department='General Trade', designation='Area Sales Manager',
                    work_location='Delhi HO', settings=self.s)
        return calculate(**{**base, **kw})

    def test_components_add_up_to_the_ctc(self):
        for ctc, cadre, grade, cat, loc in [(1200000, 'M', 'M1', 'Sales', 'Delhi HO'),
                                            (345000, 'O', 'O3', 'Plant', 'Delhi'),
                                            (300000, 'O', 'O3', 'Plant', 'Roorkee'),
                                            (3500000, 'C', 'C2', 'HO', 'Mumbai'),
                                            (180000, 'W', 'W2', 'Plant', 'Rajasthan'),
                                            (9000000, 'D', 'D1', 'HO', 'Delhi')]:
            b = self.calc(ctc_annual=ctc, cadre=cadre, grade=grade, category=cat, work_location=loc)
            self.assertNotIn('error', b, (ctc, cadre))
            gross = sum(x['monthly'] for x in b['earnings']) + sum(x['monthly'] for x in b['reimbursements'])
            self.assertEqual(gross, b['grossSalary'])
            self.assertEqual(b['grossSalary'] + b['benefitsTotal'], b['fixedMonthly'])
            self.assertEqual(b['totalCtcAnnual'], ctc, (ctc, cadre))
            self.assertEqual(sum(x['annual'] for x in b['earnings']), b['grossEarningsAnnual'])
            self.assertEqual(b['fixedAnnual'] + sum(x['annual'] for x in b['otherPayments']), ctc)

    def test_matches_p_and_c_sheet_for_14_lakh(self):
        # P & C's worked example: C1, Rs 14,00,000, two children, nothing picked.
        b = self.calc(ctc_annual=1400000, cadre='C', grade='C1', category='Plant', children=2,
                      work_location='Uttarakhand')
        e = {x['label'].split(' (')[0]: x['monthly'] for x in b['earnings']}
        ben = {x['label']: x['monthly'] for x in b['benefits']}
        self.assertEqual(b['otherPayments'][0]['monthly'], 9333)              # VPP
        self.assertEqual(b['otherPayments'][1]['monthly'], 1250)              # Gift
        self.assertEqual(ben['Mediclaim Charges'], 900)
        self.assertEqual(ben['Employer PF Contribution'], 3000)
        self.assertEqual(b['grossSalary'], 102183)
        self.assertEqual(e['Basic Salary'], 58333)                            # 50% of CTC
        self.assertEqual(e['House Rent Allowance'], 29167)                    # 50% of Basic
        self.assertEqual(e['Child Education Allowance'], 6000)
        self.assertEqual(e['Special Allowance'], 8683)
        self.assertEqual(b['grossEarnings'], 102183)
        self.assertEqual(b['totalCtcAnnual'], 1400000)
        w = self.calc(ctc_annual=180000, cadre='W', grade='W2', category='Plant', work_location='Rajasthan')
        self.assertAlmostEqual(w['earnings'][1]['monthly'], w['earnings'][0]['monthly'] * 0.2, delta=1)

    def test_vpp_and_gift_come_off_the_top(self):
        b = self.calc()
        vpp, gift = b['otherPayments']
        self.assertEqual(vpp['annual'], 120000)              # 10% of 12 lakh
        self.assertEqual(gift['monthly'], 1250)
        for cadre, grade in (('O', 'O3'), ('W', 'W2')):
            other = self.calc(ctc_annual=300000, cadre=cadre, grade=grade, category='Plant',
                              work_location='Rajasthan')['otherPayments']
            self.assertEqual(other[1]['annual'], 0, f'cadre {cadre} should get no gift')
        self.assertEqual(self.calc(category='HO', grade='M1', designation='Manager')['vppPercent'], 0)
        self.assertEqual(self.calc(category='HO', grade='M3', designation='Manager')['vppPercent'], 8)
        self.assertEqual(self.calc(category='Plant', grade='M4', designation='Manager')['vppPercent'], 8)

    def test_vpp_follows_the_variable_pay_policy(self):
        # Sales: P & C's table only - grade and department.
        for grade, dept, pct in [('C2', 'GENERAL TRADE', 15),
                                 ('M1', 'GENERAL TRADE', 10), ('M1', 'ALTERNATE CHANNELS', 12),
                                 ('M2', 'GENERAL TRADE', 10),
                                 ('M3', 'GENERAL TRADE', 10), ('M3', 'ALTERNATE CHANNELS', 12),
                                 ('M4', 'BOTH', 12), ('M4', 'GENERAL TRADE', 12), ('M4', 'ALTERNATE CHANNELS', 12),
                                 ('M5', 'ALTERNATE CHANNELS', 15)]:
            self.assertEqual(vpp_percent('Sales', grade, dept, 'Business Head')[0], pct, (grade, dept))
        # M5 and above: 15%, whatever the department.
        for grade in ('M5', 'M6', 'C1', 'C2', 'C3', 'C4', 'C5', 'D1'):
            for dept in ('GENERAL TRADE', 'ALTERNATE CHANNELS', 'BOTH', ''):
                self.assertEqual(vpp_percent('Sales', grade, dept, 'ase')[0], 15, (grade, dept))
        # Below M5, a department the table doesn't pair with the grade: the
        # grade's rate (the higher of two) - never 0.
        for grade, dept, pct in [('M2', 'ALTERNATE CHANNELS', 10), ('M1', 'BOTH', 12), ('M3', 'BOTH', 12),
                                 ('M3', 'ASM', 12), ('M1', '', 12), ('M4', '', 12)]:
            self.assertEqual(vpp_percent('Sales', grade, dept, 'Business Head')[0], pct, (grade, dept))
        self.assertEqual(vpp_percent('HO', 'C1', '', 'Manager')[0], 8)       # HO is not the Sales table
        # Managerial roles and above only.
        self.assertEqual(vpp_percent('Sales', 'O5', 'GT', 'ASM')[0], 0)
        # Corporate / Factory: 8% from M3 up ...
        for cat in ('HO', 'Plant'):
            for grade in ('M3', 'M6', 'C1', 'C5', 'D1'):
                self.assertEqual(vpp_percent(cat, grade, 'Finance', 'Manager')[0], 8, (cat, grade))
            for grade in ('M1', 'M2', 'O5', 'W1'):
                self.assertEqual(vpp_percent(cat, grade, 'Finance', 'Manager')[0], 0, (cat, grade))
        # ... except Deputy and Assistant Managers, however they're written.
        for desig in ('Deputy Manager', 'Dy. Manager - Finance', 'Assistant Manager', 'Asst. Manager (HR)'):
            pct, why = vpp_percent('HO', 'M4', 'Finance', desig)
            self.assertEqual(pct, 0, desig)
            self.assertIn('excluded', why)
        # "Deputy General Manager" is not a Deputy Manager.
        self.assertEqual(vpp_percent('Plant', 'C1', 'Ops', 'Deputy General Manager')[0], 8)

    def test_salaries_right_at_the_esi_and_pf_limits(self):
        # Where switching Mediclaim to ESI, or PF on wages to Rs 3,000, would jump
        # Gross back across the limit: ESI if within the limit with Mediclaim,
        # PF Rs 3,000 if the salary reaches its limit - and a note says so.
        pf_band = self.calc(ctc_annual=383447, cadre='C', grade='C2', category='Plant', children=1,
                            work_location='Uttar Pradesh')
        ben = {x['label']: x['monthly'] for x in pf_band['benefits']}
        self.assertEqual(ben['Employer PF Contribution'], 3000)
        self.assertTrue(any('sits right at' in n for n in pf_band['notes']))
        esi_band = self.calc(ctc_annual=326410, cadre='D', grade='D1', category='Plant',
                             work_location='Maharashtra')
        ben = {x['label']: x['monthly'] for x in esi_band['benefits']}
        self.assertGreater(ben['Employer ESI Contribution'], 0)
        self.assertEqual(ben['Mediclaim Charges'], 0)
        for b in (pf_band, esi_band):
            self.assertEqual(b['totalCtcAnnual'], b['ctcAnnual'])

    def test_esi_instead_of_mediclaim_up_to_21000(self):
        small = self.calc(ctc_annual=264000, cadre='O', grade='O5', category='HO',
                          work_location='Uttar Pradesh')
        self.assertLessEqual(small['grossSalary'], 21000)
        ben = {x['label']: x['monthly'] for x in small['benefits']}
        self.assertEqual(ben['Employer ESI Contribution'], round(small['grossSalary'] * 0.0325))
        self.assertEqual(ben['Mediclaim Charges'], 0)
        big = self.calc(ctc_annual=330000, cadre='O', grade='O3', category='Plant', work_location='Delhi')
        self.assertGreater(big['grossSalary'], 21000)
        ben = {x['label']: x['monthly'] for x in big['benefits']}
        self.assertEqual((ben['Employer ESI Contribution'], ben['Mediclaim Charges']), (0, 450))
        self.assertEqual(ben['Statutory Bonus'], 0)

    def test_pf_is_12_percent_of_gross_less_hra_below_25000_else_3000(self):
        small = self.calc(ctc_annual=288000, cadre='O', grade='O5', category='HO', work_location='Delhi')
        hra = small['earnings'][1]['monthly']
        pf = small['benefits'][0]['monthly']
        self.assertLess(small['grossSalary'], 25000)
        self.assertEqual(pf, round((small['grossSalary'] - hra) * 0.12))
        big = self.calc(ctc_annual=3500000, cadre='C', grade='C2', category='HO')
        self.assertEqual(big['benefits'][0]['monthly'], 3000)
        self.assertEqual(big['benefits'][2]['monthly'], 950)                  # mediclaim by grade

    def test_gross_earnings_answered_no_are_not_paid(self):
        base = dict(ctc_annual=3500000, cadre='C', grade='C2', category='HO', work_location='Delhi')
        full = self.calc(**base)
        some = self.calc(**base, options={'meal': True, 'telephone': False, 'car': True})
        paid = {x['key']: x['monthly'] for x in some['reimbursements']}
        self.assertEqual((paid['telephone'], paid['uniform'], paid['meal']), (0, 0, 9600))
        self.assertTrue(some['carLease'])
        self.assertFalse(full['carLease'])
        spl = lambda b: next(x['monthly'] for x in b['earnings'] if x['label'].startswith('Special'))
        # Only what is picked is paid, out of Special; Basic, HRA and PF don't move.
        self.assertEqual(sum(x['monthly'] for x in full['reimbursements']), 0)
        self.assertEqual(spl(full) - spl(some), 9600)
        self.assertEqual(some['earnings'][:2], full['earnings'][:2])
        self.assertEqual(some['benefits'], full['benefits'])
        self.assertEqual(some['totalCtcAnnual'], 3500000)
        # O and W have no Gross Earnings items, whatever is passed.
        o = self.calc(ctc_annual=600000, cadre='O', grade='O3', category='Plant',
                      work_location='Uttar Pradesh', options={'car': True})
        self.assertFalse(o['carLease'])

    def test_full_slab_reimbursements_and_special_takes_the_balance(self):
        from .rules import REIMBURSEMENTS
        everything = {k: True for k, *_ in REIMBURSEMENTS}
        b = self.calc(ctc_annual=6000000, cadre='C', grade='C2', category='HO', children=2,
                      work_location='Delhi', options=everything, tax_regime='Old Regime')
        paid = {x['key']: x['monthly'] for x in b['reimbursements']}
        lta = next(x for x in b['earnings'] if x['label'].startswith('Leave Travel'))['monthly']
        self.assertEqual({**paid, 'lta': lta}, {k: a['C'] for k, _l, a, _r in REIMBURSEMENTS})
        spl = next(x for x in b['earnings'] if x['label'].startswith('Special'))['monthly']
        self.assertGreater(spl, 0)
        others = sum(x['monthly'] for x in b['earnings'] if not x['label'].startswith('Special'))
        self.assertEqual(spl, b['grossSalary'] - others - sum(paid.values()))
        self.assertEqual(b['notes'], [])

    def test_only_what_fits_in_special_is_paid_in_preference_order(self):
        from .rules import REIMBURSEMENTS
        everything = {k: True for k, *_ in REIMBURSEMENTS}
        base = dict(ctc_annual=1400000, cadre='C', grade='C1', category='Plant', children=2,
                    work_location='Uttarakhand')
        b = self.calc(**base, options=everything)
        # Special before is 8,683: LTA 15,000 and Meal 9,600 don't fit; Telephone 5,000
        # does (3,683 left), Fuel 5,000 doesn't, Driver 3,000 does (683 left), the rest don't.
        paid = {x['key']: x['monthly'] for x in b['reimbursements'] if x['monthly']}
        self.assertEqual(paid, {'telephone': 5000, 'driver': 3000})
        self.assertEqual(b['earnings'][4]['monthly'], 683)
        self.assertTrue(any('Not enough Special Allowance' in n for n in b['notes']))
        # Books and Uniform (Old Regime only) only under Old Regime; New and Both
        # Regime get the Both Regime items only.
        rich = dict(base, ctc_annual=6000000)
        for regime, books in (('Old Regime', 5000), ('New Regime', 0), ('Both Regime', 0), ('', 0)):
            b = self.calc(**rich, options=everything, tax_regime=regime)
            paid = {x['key']: x['monthly'] for x in b['reimbursements']}
            self.assertEqual((paid['books'], paid['uniform']), (books, books), regime)
            self.assertEqual(paid['meal'], 9600, regime)                     # Both Regime item
        self.assertTrue(any('Old Regime only' in n for n in
                            self.calc(**rich, options=everything, tax_regime='Both Regime')['notes']))
        # Up to Rs 12 lakh: none at all.
        small = self.calc(**dict(base, ctc_annual=1200000), options=everything)
        self.assertEqual(sum(x['monthly'] for x in small['reimbursements']), 0)

    def test_reimbursements_are_trimmed_not_overdrawn(self):
        b = self.calc(ctc_annual=420000, category='HO', children=2, options={'meal': True})
        spl = next(x for x in b['earnings'] if x['label'].startswith('Special'))
        self.assertGreaterEqual(spl['monthly'], 0)
        self.assertTrue(any('reduced' in n for n in b['notes']))

    def test_work_location_resolves_to_a_state(self):
        for typed, state in [('Delhi HO', 'Delhi'), ('UTTAR PRADESH', 'Uttar Pradesh'),
                             ('Roorkee Plant', 'Uttarakhand'), ('Mumbai', 'Maharashtra'),
                             ('Andhra Pradesh', 'Andhra Pradesh'), ('Jammu & Kashmir', 'Jammu & Kashmir'),
                             ('Daman (UT)', 'Daman (UT)'), ('Goa', 'Goa'), ('Atlantis', '')]:
            self.assertEqual(state_for(typed), state, typed)

    def test_ctc_per_month_never_below_the_state_minimum_wage(self):
        b = self.calc(ctc_annual=345000, cadre='O', grade='O1', category='Plant', work_location='Delhi')
        self.assertNotIn('error', b)
        self.assertEqual(b['earnings'][0]['note'], '50% of CTC')
        self.assertEqual(b['earnings'][0]['monthly'], round(345000 / 12 / 2))
        unknown = self.calc(work_location='Atlantis')
        self.assertTrue(any('No minimum wage found' in n for n in unknown['notes']))
        # Checked on the CTC per month: Rs 24,000 clears Delhi's Rs 22,500 ...
        ok = self.calc(ctc_annual=288000, cadre='O', grade='O4', category='HO', work_location='Delhi')
        self.assertNotIn('error', ok)
        self.assertEqual(ok['totalCtcAnnual'], 288000)
        self.assertIsNone(ok['ctcAdjusted'])
        # ... and Rs 22,000 is raised to Delhi's Rs 22,500, with a warning.
        low = self.calc(ctc_annual=264000, cadre='O', grade='O4', category='HO', work_location='Delhi')
        self.assertNotIn('error', low)
        self.assertEqual(low['totalCtcAnnual'], 270000)
        self.assertEqual(low['ctcAdjusted'], {'enteredAnnual': 264000, 'correctedAnnual': 270000,
                                              'state': 'Delhi', 'minWage': 22500})
        self.assertIn('automatically corrected to the minimum wage of Delhi', low['notes'][0])

    def test_cea_is_3000_per_child_for_up_to_two(self):
        cea = lambda **kw: next(x for x in self.calc(**kw)['earnings'] if x['label'].startswith('Child'))
        self.assertEqual(cea(children=0)['monthly'], 0)
        self.assertEqual(cea(children=1)['monthly'], 3000)
        self.assertEqual(cea(children=2)['monthly'], 6000)
        self.assertEqual(cea(children=4)['monthly'], 6000)
        self.assertEqual(cea(ctc_annual=300000, cadre='O', grade='O3', category='Plant', children=2,
                             work_location='Rajasthan')['monthly'], 0)

    def test_unworkable_input_is_an_error_not_a_number(self):
        self.assertIn('error', self.calc(cadre='Z'))
        self.assertIn('error', self.calc(ctc_annual=10000, work_location='Atlantis'))   # no minimum wage to raise to

    def test_capped_pf_does_not_claim_a_gross_the_candidate_does_not_have(self):
        # Right at the PF limit the capped PF is paid on a gross BELOW the
        # ceiling (here Rs 24,150 against Rs 25,000). The note used to read
        # "Gross Salary Rs 25,000 or more", which was simply untrue on the
        # Annexure A the candidate is given.
        b = self.calc(ctc_annual=381000, cadre='D', grade='C1', category='HO',
                      work_location='Delhi')
        pf = next(x for x in b['benefits'] if x['label'] == 'Employer PF Contribution')
        self.assertLess(b['grossSalary'], self.s.pf_wage_ceiling)
        self.assertEqual(pf['monthly'], 3000)
        self.assertNotIn('or more', pf['note'])
        self.assertIn('PF wage ceiling', pf['note'])

    def test_the_limit_note_follows_the_ceiling_in_settings(self):
        # The capped PF is 12% of the ceiling, not a fixed Rs 3,000 -- a super
        # admin can change the ceiling, and the note has to say the real figure.
        self.s.pf_wage_ceiling = 30000
        b = self.calc(ctc_annual=452000, cadre='D', grade='C1', category='HO',
                      work_location='Delhi')
        pf = next(x for x in b['benefits'] if x['label'] == 'Employer PF Contribution')
        self.assertEqual(pf['monthly'], 3600)
        note = next((n for n in b['notes'] if 'right at the' in n), '')
        self.assertIn('Rs 3,600', note)
        self.assertNotIn('Rs 3,000', note)


class CtcAccessIsGranted(TestCase):
    """The Admin Console matrix has to be a control, not a courtesy.

    Candidate names and the salaries they are offered sit behind these
    endpoints. Before the CTC Calculator was a grantable tool every signed-in
    employee -- all 664 of them -- could open it, because the only check was
    "are you signed in" and the screen decided the rest.
    """
    def setUp(self):
        self.granted = PortalUser.objects.create(email='pc@apisindia.com', employee_code='P1',
                                                 name='P and C', app_access=[AppKey.CTC])
        self.plain = PortalUser.objects.create(email='emp@apisindia.com', employee_code='P2',
                                               name='Employee')
        self.admin = PortalUser.objects.create(email='sa@apisindia.com', employee_code='P3',
                                               name='Super', is_superadmin=True)

    def _auth(self, user):
        return {'HTTP_AUTHORIZATION': f'Bearer {PortalSession.start(user)}'}

    def test_an_employee_without_the_grant_is_refused_everywhere(self):
        for path in ('/template/', '/uploads/', '/settings/'):
            r = self.client.get(f'{BASE}{path}', **self._auth(self.plain))
            self.assertEqual(r.status_code, 403, path)
        r = self.client.post(f'{BASE}/uploads/', {}, **self._auth(self.plain))
        self.assertEqual(r.status_code, 403)

    def test_the_grant_opens_it_and_a_super_admin_never_needs_one(self):
        for user in (self.granted, self.admin):
            for path in ('/template/', '/uploads/', '/settings/'):
                r = self.client.get(f'{BASE}{path}', **self._auth(user))
                self.assertEqual(r.status_code, 200, (user.email, path))

    def test_signed_out_is_401_not_403(self):
        # 401 is what tells the browser the token is dead; 403 would leave a
        # stale token in place and the person stuck on a blank screen.
        self.assertEqual(self.client.get(f'{BASE}/uploads/').status_code, 401)

    def test_the_tool_is_in_the_admin_console_matrix(self):
        # The column only exists because AppKey does -- that is what makes a
        # new tool controllable without anyone remembering to wire it up.
        self.assertIn('ctc-calculator', AppKey.values)
        self.assertEqual(AppKey.CTC.label, 'CTC Calculator')
