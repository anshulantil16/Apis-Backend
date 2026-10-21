"""The CTC Calculator template: download, fill in one row per candidate,
upload back on the same page.

Generated rather than kept as a file, so it always matches what the
uploader expects (see columns.py). The first sheet holds nothing but the
header row — an example row there would be uploaded by mistake sooner or
later — and the help sits on the header cells as comments, with a worked
example on a second sheet.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .columns import CHOICES, COLUMNS, OPTION_COLUMNS, SALES_DEPARTMENTS, TAX_REGIMES
from .rules import MIN_WAGES, REIMBURSEMENT_MIN_CTC, REIMBURSEMENTS

SHEET_NAME = 'Candidates'

EXAMPLE = ['Rahul Sharma', 'Sales', 'GENERAL TRADE', 'Area Sales Manager', 'M1', 'M', 2,
           'Delhi', 1500000, 'New Regime', 'Yes', 'Yes', 'Yes', 'No', 'No', 'Yes', 'No', 'No',
           'No']

HEADER_FILL = PatternFill('solid', fgColor='F59E0B')     # APIS amber
REQUIRED_FILL = PatternFill('solid', fgColor='D97706')
HEADER_FONT = Font(bold=True, color='FFFFFF')
OPTION_FILL = PatternFill('solid', fgColor='00B0F0')      # the Annexure A's blue
DISABLED_FILL = PatternFill('solid', fgColor='D9D9D9', bgColor='D9D9D9')
OPTION_WIDTH = 22


def _col(field):
    return next(get_column_letter(i) for i, (_h, f, *_r) in enumerate(COLUMNS, start=1) if f == field)


def _option_columns(ws, lists):
    """The Gross Earnings Yes / No columns after CTC, and the Tax Regime.
    On a row whose Cadre is O or W, or whose CTC is Rs 12 lakh or less, the
    Gross Earnings cells are greyed out and their dropdown is empty - there
    is nothing to pick, and anything typed is refused; the Tax Regime likewise
    for a CTC of Rs 12 lakh or less."""
    cadre, ctc, regime = _col('cadre'), _col('ctc'), _col('tax_regime')
    # Each row reads its own Cadre and CTC through INDEX(column, ROW()) rather
    # than $I2: Excel rewrites a plain reference when cells are cut and pasted
    # or dragged, which left rows checking a neighbour's CTC.
    own = lambda col: f'INDEX(${col}:${col},ROW())'
    small = f'N({own(ctc)})<={REIMBURSEMENT_MIN_CTC}'
    regimes = {k: r for k, _l, _a, r in REIMBURSEMENTS}
    first = len(COLUMNS) + 1
    last = len(COLUMNS) + len(OPTION_COLUMNS)
    for i, (heading, key, _blank) in enumerate(OPTION_COLUMNS, start=first):
        cell = ws.cell(row=1, column=i, value=heading)
        cell.font = HEADER_FONT
        cell.fill = OPTION_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.comment = Comment(
            'Gross Earnings - Yes or No. Blank means No.\n\nYes pays it at the cadre slab out of the '
            'Special Allowance left, in preference order - only if it fits in full.'
            + ('\n\nOld Regime only - not given under New Regime or Both Regime.'
               if regimes.get(key) == 'old' else '')
            + '\n\nOnly for cadres D, C and M with a CTC above Rs 12 lakh - greyed out otherwise.',
            'APIS Intranet')
        ws.column_dimensions[get_column_letter(i)].width = OPTION_WIDTH
    ws.row_dimensions[1].height = 48
    lists['B1'], lists['B2'] = 'Yes', 'No'          # Lists!C1 stays empty: the disabled list
    for r, name in enumerate(TAX_REGIMES, start=1):
        lists.cell(row=r, column=4, value=name)
    for r, name in enumerate(SALES_DEPARTMENTS, start=1):
        lists.cell(row=r, column=5, value=name)

    # Department: on a Sales row, a dropdown of the three sales departments;
    # otherwise an empty list, so HO and Plant type their own. (Not enforced
    # here - the upload checks a Sales row's department.)
    dept, category = _col('department'), _col('category')
    dd = DataValidation(type='list',
                        formula1=f'IF({own(category)}="Sales",Lists!$E$1:$E${len(SALES_DEPARTMENTS)},Lists!$C$1)',
                        allow_blank=True, showErrorMessage=False, showInputMessage=True,
                        promptTitle='Department',
                        prompt='Sales: pick from the list.\nHO / Plant: just type it (the list is empty).')
    dd.add(f'{dept}2:{dept}2001')
    ws.add_data_validation(dd)

    a, b = get_column_letter(first), get_column_letter(last)
    off = f'OR({own(cadre)}="O",{own(cadre)}="W",{small})'
    dv = DataValidation(type='list', formula1=f'IF({off},Lists!$C$1,Lists!$B$1:$B$2)',
                        allow_blank=True, showErrorMessage=True,
                        errorTitle='Gross Earnings',
                        error='Pick Yes or No. Only for cadres D, C and M with a CTC above '
                              'Rs 12 lakh - otherwise leave it blank.')
    dv.add(f'{a}2:{b}2001')
    ws.add_data_validation(dv)
    ws.conditional_formatting.add(f'{a}2:{b}2001', FormulaRule(formula=[off], fill=DISABLED_FILL))

    rv = DataValidation(type='list',
                        formula1=f'IF({small},Lists!$C$1,Lists!$D$1:$D${len(TAX_REGIMES)})',
                        allow_blank=True, showErrorMessage=True, errorTitle='Tax Regime',
                        error='Pick Old Regime, New Regime or Both Regime. Only for a CTC above '
                              'Rs 12 lakh - otherwise leave it blank.')
    rv.add(f'{regime}2:{regime}2001')
    ws.add_data_validation(rv)
    ws.conditional_formatting.add(f'{regime}2:{regime}2001', FormulaRule(formula=[small], fill=DISABLED_FILL))


def build_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    for i, (heading, _field, width, required, help_text) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i, value=heading)
        cell.font = HEADER_FONT
        cell.fill = REQUIRED_FILL if required else HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.comment = Comment(help_text, 'APIS Intranet')
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = 'A2'
    lists = wb.create_sheet('Lists')
    lists.sheet_state = 'hidden'
    _option_columns(ws, lists)

    # Dropdowns, so Category, Cadre and Number of Children can only be typed one way.
    for i, (_h, field, *_rest) in enumerate(COLUMNS, start=1):
        if field in CHOICES:
            dv = DataValidation(type='list', formula1='"' + ','.join(CHOICES[field]) + '"',
                                allow_blank=True, showErrorMessage=True,
                                errorTitle='Pick from the list',
                                error='Choose one of: ' + ', '.join(CHOICES[field]))
            col = get_column_letter(i)
            dv.add(f'{col}2:{col}2001')
            ws.add_data_validation(dv)

    # Work Location: a dropdown of states, kept on a hidden sheet because the
    # list is longer than Excel allows inline. A suggestion, not a rule — a
    # city typed instead still uploads and is matched to its state.
    for r, state in enumerate(MIN_WAGES, start=1):
        lists.cell(row=r, column=1, value=state)
    loc_col = next(get_column_letter(i) for i, (_h, f, *_r) in enumerate(COLUMNS, start=1)
                   if f == 'work_location')
    loc = DataValidation(type='list', formula1=f'Lists!$A$1:$A${len(MIN_WAGES)}',
                         allow_blank=True, showErrorMessage=False)
    loc.add(f'{loc_col}2:{loc_col}2001')
    ws.add_data_validation(loc)

    ex = wb.create_sheet('Example')
    ex.cell(row=1, column=1, value='How a filled-in row looks. This sheet is never uploaded — '
                                   'type your candidates on the "Candidates" sheet.').font = Font(italic=True, color='64748B')
    headings = [(h, w, HEADER_FILL) for h, _f, w, *_ in COLUMNS] + \
               [(h, OPTION_WIDTH, OPTION_FILL) for h, *_ in OPTION_COLUMNS]
    for i, (heading, width, fill) in enumerate(headings, start=1):
        h = ex.cell(row=3, column=i, value=heading)
        h.font = HEADER_FONT
        h.fill = fill
        h.alignment = Alignment(wrap_text=True, vertical='center')
        ex.cell(row=4, column=i, value=EXAMPLE[i - 1])
        ex.column_dimensions[get_column_letter(i)].width = width
    ex.row_dimensions[3].height = 48

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
