"""Read a filled-in CTC template into candidate rows.

Headings are matched loosely — case, spacing and stray punctuation are
ignored — because "candidate name ", "Candidate Name" and "CANDIDATE_NAME"
all mean the same thing to the person typing them. Only the first sheet, or
the one named "Candidates", is read.

Nothing is saved here. parse() returns the rows and the problems it found,
and the view decides: any problem and nothing is stored, so a half-imported
sheet never has to be untangled.
"""
import re
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from .calculator import normalise_cadre, normalise_category, sales_department
from .rules import REIMBURSEMENT_MIN_CTC
from .columns import (COLUMNS, MAX_CTC, MAX_LENGTH, NO, NOT_ANSWERED, OPTION_CADRES,
                      OPTION_COLUMNS, OPTIONAL_FIELDS, TAX_REGIMES, YES)
from .template import SHEET_NAME

MAX_ROWS = 2000


class SheetError(Exception):
    """The workbook as a whole can't be read. Message is for the uploader."""


def _key(text):
    return re.sub(r'[^a-z0-9]', '', str(text or '').lower())


def _clean(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r'\s+', ' ', str(value)).strip()


def _amount(text):
    """A rupee amount typed any reasonable way — 1200000, 12,00,000,
    ₹ 12,00,000, 1200000.50 — as a Decimal, or None if it isn't one."""
    cleaned = re.sub(r'(?i)(rs\.?|inr|₹|,|\s)', '', text)
    if not cleaned:
        return None
    try:
        value = Decimal(cleaned)
    except InvalidOperation:
        return None
    return value if value.is_finite() else None


def _regime(text):
    """"old", "Old Regime", "NEW" ... -> the template's spelling, or ''."""
    k = _key(text).replace('regime', '').replace('regim', '')
    return next((r for r in TAX_REGIMES if _key(r).replace('regime', '') == k), '')


def parse(fileobj):
    """(rows, errors). Each row is {'row_number': n, <field>: value, ...}."""
    try:
        wb = load_workbook(fileobj, read_only=True, data_only=True)
    except Exception:
        raise SheetError("That file couldn't be read as an Excel workbook. "
                         'Please use the downloaded template (.xlsx).')
    try:
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[0]
        lines = ws.iter_rows(values_only=True)
        header = next(lines, None)
        if not header:
            raise SheetError('The sheet is empty. Fill in the template and upload it again.')

        wanted = {_key(heading): field for heading, field, *_ in COLUMNS}
        position = {}
        for i, h in enumerate(header):
            field = wanted.get(_key(h))
            if field and field not in position:
                position[field] = i
        options = {}                                     # the optional Gross Earnings columns
        for i, h in enumerate(header):
            for heading, key, _blank in OPTION_COLUMNS:
                if _key(h) == _key(heading) and key not in options:
                    options[key] = i
        missing = [heading for heading, field, *_ in COLUMNS
                   if field not in position and field not in OPTIONAL_FIELDS]
        if missing:
            raise SheetError('These columns are missing from the first row: '
                             + ', '.join(missing) + '. Please use the downloaded template.')

        rows, errors = [], []
        for n, values in enumerate(lines, start=2):
            values = values or ()
            raw = {field: (values[i] if i < len(values) else None) for field, i in position.items()}
            row = {field: _clean(v) for field, v in raw.items()}
            for field in OPTIONAL_FIELDS:
                row.setdefault(field, '')
            picked = {key: _clean(values[i] if i < len(values) else None)
                      for key, i in options.items()}
            if not any(row.values()):
                continue                                  # blank line
            if len(rows) >= MAX_ROWS:
                raise SheetError(f'Please upload at most {MAX_ROWS} candidates at a time.')
            for heading, field, _w, required, _help in COLUMNS:
                if required and field != 'ctc' and not row[field]:
                    errors.append(f'Row {n}: {heading} is empty.')
            if row['category']:
                category = normalise_category(row['category'])
                if category:
                    row['category'] = category
                else:
                    errors.append(f'Row {n}: Category "{row["category"]}" should be Sales, HO or Plant.')
            if row['category'] == 'Sales':
                dept = sales_department(row['department'])
                if dept:
                    row['department'] = dept
                else:
                    errors.append(f'Row {n}: Department "{row["department"]}" - for Sales pick GENERAL TRADE, '
                                  'ALTERNATE CHANNELS or BOTH (it sets the Variable Pay %).')
            if row['tax_regime'].strip().lower() in NOT_ANSWERED:
                row['tax_regime'] = ''                     # 'N/A', '-' ... from an older template
            if row['tax_regime']:
                regime = _regime(row['tax_regime'])
                if regime:
                    row['tax_regime'] = regime
                else:
                    errors.append(f'Row {n}: Tax Regime "{row["tax_regime"]}" should be '
                                  'Old Regime or New Regime.')
            if row['cadre']:
                cadre = normalise_cadre(row['cadre'])
                if cadre:
                    row['cadre'] = cadre
                else:
                    errors.append(f'Row {n}: Cadre "{row["cadre"]}" should be D, C, M, O or W.')
            if row['grade']:
                row['grade'] = row['grade'].upper()
                if row['cadre'] in 'DCMOW' and row['cadre'] and not row['grade'].startswith(row['cadre']):
                    errors.append(f'Row {n}: Grade {row["grade"]} does not belong to cadre {row["cadre"]}.')
            ctc_text, row['ctc'] = row['ctc'], None
            if not ctc_text:
                errors.append(f'Row {n}: CTC is empty.')
            else:
                amount = _amount(ctc_text)
                if amount is None:
                    errors.append(f'Row {n}: CTC "{ctc_text}" is not a number.')
                elif amount <= 0 or amount > MAX_CTC:
                    errors.append(f'Row {n}: CTC {ctc_text} is outside the expected range.')
                else:
                    row['ctc'] = amount.quantize(Decimal('0.01'))
            kids_text, row['number_of_children'] = row['number_of_children'], 0
            if kids_text:
                if kids_text.isdigit() and int(kids_text) <= 2:
                    row['number_of_children'] = int(kids_text)
                else:
                    errors.append(f'Row {n}: Number of Children "{kids_text}" should be 0, 1 or 2.')
            for field, value in row.items():
                if field in ('ctc', 'number_of_children', 'gross_earnings'):
                    continue
                if len(value) > MAX_LENGTH[field]:
                    heading = next(h for h, f, *_ in COLUMNS if f == field)
                    errors.append(f'Row {n}: {heading} is longer than {MAX_LENGTH[field]} characters.')
            row['gross_earnings'] = {}
            if row['cadre'] in OPTION_CADRES:            # O and W: greyed out, ignored
                for heading, key, _blank in OPTION_COLUMNS:
                    answer = picked.get(key, '').lower()
                    if answer in YES:
                        row['gross_earnings'][key] = True
                    elif answer in NO:
                        row['gross_earnings'][key] = False
                    elif answer not in NOT_ANSWERED:
                        errors.append(f'Row {n}: {heading} "{picked[key]}" should be Yes or No.')
            if row['ctc'] is not None and row['ctc'] <= REIMBURSEMENT_MIN_CTC:
                row['tax_regime'], row['gross_earnings'] = '', {}   # greyed out: ignored
            rows.append({'row_number': n, **row})
        return rows, errors
    finally:
        wb.close()
