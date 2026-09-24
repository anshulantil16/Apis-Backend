"""A month's work as a workbook, for handing to a manager.

Two sheets, because a manager asks two things of the same data: how much did
each person do, and what was it. The second sheet is one row per job, flat,
with the person repeated on every row -- a shape that can be filtered and
pivoted, which is the only reason to send a spreadsheet rather than a number.

Both kinds of work are in it and the column that says which is not optional.
"Thirty jobs" means something different depending on whether people are
raising tickets or the team is writing its own work up, and a manager reading
this should be able to see that for themselves.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill('solid', fgColor='0E7490')
THIN = Side(style='thin', color='E2E8F0')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SUMMARY_COLUMNS = [
    ('Person', 26), ('Email', 30), ('From tickets', 14),
    ('Logged directly', 16), ('Total jobs', 12), ('Time recorded', 15),
]
DETAIL_COLUMNS = [
    ('Date', 12), ('Person', 24), ('How it came in', 16), ('Category', 26),
    ('What was done', 46), ('For', 24), ('Minutes', 10), ('State', 14),
]


def _head(ws, columns):
    for i, (label, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i, value=label)
        cell.font = Font(bold=True, color='FFFFFF', size=10)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = 'A2'


def _row(ws, r, values):
    for i, v in enumerate(values, start=1):
        cell = ws.cell(row=r, column=i, value=v)
        cell.border = BORDER
        cell.alignment = Alignment(vertical='top',
                                   wrap_text=isinstance(v, str) and len(v) > 40)


def _minutes(m):
    if not m:
        return ''
    return f'{m // 60}h {m % 60}m' if m >= 60 else f'{m}m'


def build_work_report(report, tickets):
    """-> xlsx bytes. `report` is WorkReportView's dict, `tickets` the rows."""
    wb = Workbook()

    ws = wb.active
    ws.title = 'Summary'
    ws['A1'] = f"Work done — {report['label']}"
    ws['A1'].font = Font(bold=True, size=13)
    ws['A3'] = (f"{report['total']} jobs: {report['from_tickets']} from tickets, "
                f"{report['logged_directly']} logged directly by the team.")
    ws['A3'].font = Font(size=10, color='475569')
    ws['A4'] = ('Counted by the day the work was done. A ticket raised last month and '
                'finished this one belongs to this month.')
    ws['A4'].font = Font(size=9, italic=True, color='94A3B8')

    start = 6
    for i, (label, width) in enumerate(SUMMARY_COLUMNS, start=1):
        cell = ws.cell(row=start, column=i, value=label)
        cell.font = Font(bold=True, color='FFFFFF', size=10)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal='center')
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = width

    r = start + 1
    for p in report['people']:
        _row(ws, r, [p['name'], p['email'], p['closed'], p['logged'],
                     p['total'], _minutes(p['minutes'])])
        r += 1
    if not report['people']:
        ws.cell(row=r, column=1, value='Nothing recorded for this month.').font = \
            Font(italic=True, color='94A3B8')

    detail = wb.create_sheet('Every job')
    _head(detail, DETAIL_COLUMNS)
    for i, t in enumerate(tickets, start=2):
        _row(detail, i, [
            t.performed_on.strftime('%d %b %Y') if t.performed_on else '',
            t.performed_by_name or t.performed_by_email or 'Unattributed',
            'Logged' if t.origin == 'logged' else 'From a ticket',
            t.get_category_display(),
            t.subject,
            t.logged_for or t.requested_by_name or '',
            t.time_spent_minutes or '',
            t.get_status_display(),
        ])
    if not tickets:
        detail.cell(row=2, column=1, value='No jobs recorded for this month.').font = \
            Font(italic=True, color='94A3B8')

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
