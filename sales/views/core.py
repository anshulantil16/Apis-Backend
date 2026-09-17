"""SalesIQ core endpoints — upload, template, KPIs, breakdowns, trend,
forecast, insights, uploads management and Excel export."""
import io
from datetime import date, timedelta

import openpyxl
from django.db.models import Count, Max, Min, Sum
from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.parsers import MultiPartParser, FormParser

from ..models import SalesUpload, SalesRecord, sync_actual_source
from ..ingest import (map_headers, parse_date, parse_num, build_template,
                      TEXT_FIELDS, NUM_FIELDS, TEXT_MAX, DATE_FIELDS,
                      parse_bool, is_return_type, partition_unknown,
                      state_from_code, state_from_subregion, find_header_row)
from .. import aop as AOP

from ..forecasting import forecast_series
from .filters import (DIMENSIONS, FILTERABLE, _multi, apply_filters,
                      apply_dim_filters, _period_bounds, _money, _pct_change)

class SalesTemplateView(APIView):
    def get(self, request):
        buf = build_template()
        resp = HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp['Content-Disposition'] = 'attachment; filename="SalesIQ_Template.xlsx"'
        return resp


def _ingest_aop(request, upload, ws, header_row, header_row_index=1):
    """Load the AOP-vs-ACH sheet, unpivoting each row into one per month."""
    dims, months, unknown = AOP.map_columns(header_row)

    if not months:
        return Response({'error': 'No month columns found. This sheet should carry '
                                  'columns like "Apr-26 AOP" and "Apr-26".'}, status=400)

    # One upload row per FILE, created by the caller. A sheet does not get
    # its own: a two-tab workbook was appearing in the list twice under the
    # same name, which reads as having uploaded it twice.
    watermark = SalesRecord.objects.aggregate(m=Max('id'))['m'] or 0

    batch, total_rev, total_target = [], 0.0, 0.0
    lo = hi = None
    empty_rows = 0
    row_no = header_row_index
    try:
        for row in ws.iter_rows(min_row=header_row_index + 1, values_only=True):
            row_no += 1
            if not any(v is not None and str(v).strip() != '' for v in row):
                continue
            entries = AOP.unpivot(row, dims, months)
            if not entries:
                empty_rows += 1
                continue
            for e in entries:
                period = e['period']
                rec = SalesRecord(
                    upload=upload,
                    source=SalesRecord.SOURCE_PLAN,
                    order_date=period,
                    period=period,
                    net_amount=e['net_amount'],
                    plan_achievement=e['net_amount'],
                    target_amount=e['target_amount'],
                    sfo_count=e['sfo_count'],
                )
                for field in ('channel', 'sales_head', 'rsm', 'asm', 'zone', 'state',
                              'item_alt_code', 'product_name', 'brand'):
                    setattr(rec, field, str(e.get(field) or '')[:TEXT_MAX.get(field, 200)])
                # Sub-Region is a selling territory (AP-1), not a state. It
                # keeps its own column, and the state is read off its code
                # so this sheet and the invoice dump name states the same way.
                rec.subzone = rec.state[:100]
                rec.state = state_from_subregion(rec.subzone)
                batch.append(rec)
                total_rev += e['net_amount']
                total_target += e['target_amount']
                lo = period if lo is None or period < lo else lo
                hi = period if hi is None or period > hi else hi

                if len(batch) >= 2000:
                    SalesRecord.objects.bulk_create(batch, batch_size=1000)
                    batch = []
        if batch:
            SalesRecord.objects.bulk_create(batch, batch_size=1000)
    except Exception as e:
        # Only this sheet's rows. The upload belongs to the file, and another
        # sheet may already have loaded cleanly into it.
        upload.records.filter(id__gt=watermark).delete()
        return Response({'error': f'Failed while reading row {row_no}: {e}'}, status=400)

    count = upload.records.filter(id__gt=watermark).count()
    if count == 0:
        return Response({'error': 'No usable rows — every row was empty across all months.'},
                        status=400)

    warnings, notes = [], []
    plan_months = sorted({m for _, m, t in months if t})
    notes.append(
        f'{count:,} month-rows built from this sheet — each row was spread across the '
        f'{len(set(m for _, m, _ in months))} month columns it carries.')
    if plan_months:
        notes.append(
            f'Plan (AOP) loaded for {len(plan_months)} months, '
            f'{plan_months[0]:%b %Y} to {plan_months[-1]:%b %Y}. '
            f'Targets and achievement now sit side by side on the dashboard.')
    skipped, unrecognised = [], []
    for h in unknown:
        (skipped if h in AOP.IGNORED else unrecognised).append(h)
    if skipped:
        notes.append(
            f'{len(skipped)} total column(s) skipped on purpose: {", ".join(skipped)}. '
            f'Each is its own monthly columns added up, and storing a total beside the '
            f'parts it is made of double-counts the year.')
    if unrecognised:
        warnings.append(
            f'{len(unrecognised)} column(s) were not recognised and are ignored: '
            f'{", ".join(unrecognised[:12])}.')
    if empty_rows:
        notes.append(f'{empty_rows} row(s) had no figure in any month and were skipped.')

    # The reason both files can live in one table without inflating anything.
    if SalesRecord.objects.filter(source=SalesRecord.SOURCE_INVOICE).exists():
        notes.append(
            'Invoice-level data is already loaded, so sales figures keep coming from '
            'it and this sheet supplies the targets. Its own achievement columns are '
            'stored for reconciliation rather than added on top.')


    return Response({
        'message': f'Imported {count:,} month-rows.',
        'upload_id': upload.id,
        'rows': count,
        'skipped': empty_rows,
        'total_revenue': _money(total_rev),
        'total_target': _money(total_target),
        'file_kind': 'aop',
        'period': {'from': lo.isoformat() if lo else None,
                   'to': hi.isoformat() if hi else None},
        'detected_columns': sorted(dims.keys()),
        'skipped_columns': skipped,
        'unrecognised_columns': unrecognised,
        'cancelled_rows': 0,
        'return_rows': 0,
        'warnings': warnings,
        'notes': notes,
    })


class SalesUploadView(APIView):
    parser_classes = (MultiPartParser, FormParser)

    def post(self, request):
        """Load every sheet in the workbook, routing each by its own headers.

        The two primary-sales files are normally kept as two tabs of one
        workbook. Reading only wb.active meant whichever tab happened to be
        selected when the file was last saved decided what got imported — the
        other one was dropped without a word, and if the selected tab was the
        AOP sheet the dump parser rejected it with a complaint about a missing
        date column that was really a complaint about the wrong sheet.
        """
        f = request.FILES.get('file')
        if not f:
            return Response({'error': 'No file provided.'}, status=400)
        try:
            wb = openpyxl.load_workbook(f, data_only=True, read_only=True)
        except Exception as e:
            return Response({'error': f'Cannot read file: {e}'}, status=400)

        # One row in the uploaded-files list per FILE. Each sheet used to
        # create its own, so a two-tab workbook appeared twice under the same
        # name — which reads as having uploaded it twice by mistake.
        upload = SalesUpload.objects.create(
            filename=(f.name or '')[:255],
            uploaded_by=str(request.query_params.get('user') or '')[:200],
            status='completed',
        )

        outcomes = []
        for ws in wb.worksheets:
            header_row, header_at = find_header_row(
                ws, recognises=lambda r: max(len(map_headers(r)[0]),
                                             len(AOP.map_columns(r)[0])
                                             + len(AOP.map_columns(r)[1])))
            if header_row is None:
                continue        # an empty tab is not an error
            if AOP.looks_like_aop_sheet(header_row):
                resp = _ingest_aop(request, upload, ws, header_row, header_at)
            else:
                resp = _ingest_dump(request, upload, ws, header_row, header_at)
            outcomes.append((ws.title, resp))

        if not outcomes:
            upload.delete()
            return Response({'error': 'The workbook has no sheets with any data in them.'},
                            status=400)

        good = [(t, r) for t, r in outcomes if r.status_code == 200]
        if not good:
            # Nothing loaded, so nothing to show. An upload row left behind
            # here is the "0 rows - Rs 0" ghost that used to sit in the list
            # for good, because failing only ever set a status and deleted
            # the records underneath it.
            upload.delete()
            # Every sheet failed. One sheet is the common case, so answer with
            # its own error rather than a summary nobody can act on.
            title, resp = outcomes[0]
            if len(outcomes) > 1:
                resp.data['error'] = (
                    f'Nothing could be imported. "{title}": '
                    f'{resp.data.get("error", "unreadable")}')
                resp.data['sheets'] = {t: r.data.get('error') for t, r in outcomes}
            return resp

        # Counted back off the rows that actually landed, rather than carried
        # along as the sheets are read. A running total and the table it
        # describes drift the moment anything is skipped, and the list header
        # then disagrees with the row beneath it.
        sync_actual_source()
        agg = upload.records.aggregate(
            n=Count('id'), lo=Min('order_date'), hi=Max('order_date'))
        # Revenue is counted over what actually sold. A cancelled invoice is
        # stored and reported in the row count — it is a real document — but
        # it is not money, here any more than on the dashboard.
        earned = upload.records.exclude(is_cancelled=True).aggregate(
            rev=Sum('net_amount'))['rev']
        upload.row_count = agg['n'] or 0
        upload.skipped_rows = sum(r.data.get('skipped', 0) for _, r in good)
        upload.total_revenue = round(float(earned or 0), 2)
        upload.period_start, upload.period_end = agg['lo'], agg['hi']
        upload.warnings = [f'{t}: {w}' for t, r in good for w in r.data.get('warnings', [])]
        upload.notes = [f'{t}: {n}' for t, r in good for n in r.data.get('notes', [])]
        for title, resp in outcomes:
            if resp.status_code != 200:
                upload.warnings.insert(
                    0, f'{title}: not imported - {resp.data.get("error", "unreadable")}')
        upload.save()

        if len(good) == 1 and len(outcomes) == 1:
            single = good[0][1]
            single.data['rows'] = upload.row_count
            single.data['total_revenue'] = _money(upload.total_revenue)
            return single

        # More than one sheet loaded: merge into one answer, and say which
        # sheet each figure came from.
        merged = {
            'message': f'Imported {upload.row_count:,} rows from {len(good)} sheet(s).',
            'rows': upload.row_count,
            'skipped': sum(r.data.get('skipped', 0) for _, r in good),
            'total_revenue': _money(upload.total_revenue),
            'cancelled_rows': sum(r.data.get('cancelled_rows', 0) for _, r in good),
            'return_rows': sum(r.data.get('return_rows', 0) for _, r in good),
            'sheets': {t: {'rows': r.data.get('rows', 0),
                           'kind': r.data.get('file_kind', 'dump')} for t, r in good},
            'detected_columns': sorted({c for _, r in good
                                        for c in r.data.get('detected_columns', [])}),
            'skipped_columns': sorted({c for _, r in good
                                       for c in r.data.get('skipped_columns', [])}),
            'unrecognised_columns': sorted({c for _, r in good
                                            for c in r.data.get('unrecognised_columns', [])}),
            'warnings': upload.warnings,
            'notes': upload.notes,
        }
        return Response(merged)


def _ingest_dump(request, upload, ws, header_row, header_row_index=1):
    """Load one Pre-Sales Dump sheet."""
    col_map, unknown = map_headers(header_row)
    # The Pre-Sales Dump's sales value is Taxable Amount, which is its own
    # column rather than a spelling of "Net Amount" — so it counts here.
    VALUE_COLUMNS = ('net_amount', 'gross_amount', 'taxable_amount')
    if not any(c in col_map for c in ('order_date', 'invoice_date', 'posting_date')):
        return Response({
            'error': 'No date column found. One column must be the order/invoice date.',
            'detected_columns': sorted(col_map.keys()),
            'unrecognised_columns': unknown,
        }, status=400)
    if not any(c in col_map for c in VALUE_COLUMNS):
        return Response({
            'error': 'No sales value column found. Add a "Net Amount" '
                     '(or "Amount" / "Sales Value") column.',
            'detected_columns': sorted(col_map.keys()),
            'unrecognised_columns': unknown,
        }, status=400)

    def cell(row, field):
        ci = col_map.get(field)
        if ci is None or ci >= len(row):
            return None
        return row[ci]

    # One upload row per FILE, created by the caller. A sheet does not get
    # its own: a two-tab workbook was appearing in the list twice under the
    # same name, which reads as having uploaded it twice.
    watermark = SalesRecord.objects.aggregate(m=Max('id'))['m'] or 0

    batch, total_rev = [], 0.0
    no_date = bad_value = 0
    cancelled_rows = return_rows = 0
    # What the Type column actually holds, and how much money sits on
    # lines that carry no product.
    line_types = {}
    non_item_value = 0.0
    return_negative = return_positive = 0
    lo = hi = None
    row_no = header_row_index
    try:
        for row in ws.iter_rows(min_row=header_row_index + 1, values_only=True):
            row_no += 1
            if not any(v is not None and str(v).strip() != '' for v in row):
                continue
            # Invoice Date and Posting Date now map to their own columns,
            # so they no longer double as aliases for Order Date. An export
            # that carries only one of them must still land in a month, so
            # they remain the fallback here, in that order.
            od = (parse_date(cell(row, 'order_date'))
                  or parse_date(cell(row, 'invoice_date'))
                  or parse_date(cell(row, 'posting_date')))
            if od is None:
                no_date += 1
                continue

            # ── is this row a sale at all? ────────────────────────
            # An ERP dump carries cancelled invoices and credit memos in
            # the same sheet as live sales. Both are kept — they are real
            # documents and the ledger has to reconcile — but a cancelled
            # row must never reach a sales figure.
            cancelled = parse_bool(cell(row, 'is_cancelled'))
            doc_type = cell(row, 'document_type')
            returned = is_return_type(doc_type)
            if cancelled:
                cancelled_rows += 1
            if returned:
                return_rows += 1
            lt = str(cell(row, 'line_type') or '').strip()
            if lt:
                line_types[lt] = line_types.get(lt, 0) + 1

            # ── money ─────────────────────────────────────────────────
            # Taxable Amount is the sales value the business reports on:
            # after scheme and discount, before GST. It wins over a
            # generic "amount" column when the dump carries both.
            taxable = parse_num(cell(row, 'taxable_amount'))
            gross = parse_num(cell(row, 'gross_amount'))
            net = parse_num(cell(row, 'net_amount'))

            inv_disc = parse_num(cell(row, 'invoice_discount'))
            retail = parse_num(cell(row, 'retail_scheme'))
            wholesale = parse_num(cell(row, 'wholesale_scheme'))
            igst = parse_num(cell(row, 'igst_amount'))
            cgst = parse_num(cell(row, 'cgst_amount'))
            sgst = parse_num(cell(row, 'sgst_amount'))

            if taxable:
                net = taxable
            # Discount and tax are summed from their parts when the dump
            # splits them, rather than asking for a pre-totalled column
            # that would then disagree with the parts beside it.
            discount = parse_num(cell(row, 'discount')) or (inv_disc + retail + wholesale)
            tax = parse_num(cell(row, 'tax')) or (igst + cgst + sgst)

            if not gross and net:
                gross = net + discount
            # Fall back to gross when the export has no explicit net column.
            if not net and gross:
                net = gross - discount
            if not net and not gross and not cancelled:
                bad_value += 1

            if lt and lt.lower() != 'item':
                non_item_value += net
            if returned:
                if net < 0:
                    return_negative += 1
                elif net > 0:
                    return_positive += 1

            rec = SalesRecord(
                upload=upload,
                order_date=od,
                period=od.replace(day=1),
                quantity=parse_num(cell(row, 'quantity')),
                unit_price=parse_num(cell(row, 'unit_price')),
                gross_amount=gross,
                discount=discount,
                tax=tax,
                net_amount=net,
                target_amount=parse_num(cell(row, 'target_amount')),
                is_cancelled=cancelled,
                is_return=returned,
            )
            for tf in TEXT_FIELDS:
                v = cell(row, tf)
                s = '' if v is None else str(v).strip()
                setattr(rec, tf, s[:TEXT_MAX.get(tf, 150)])
            # The dump names no state, only a GST code. Filled in only
            # when the sheet gave no state of its own, so a file that does
            # carry one keeps its own spelling.
            if not rec.state:
                rec.state = state_from_code(rec.customer_state_code)
            for nf in NUM_FIELDS:
                # The headline five are derived above; the rest are stored
                # verbatim so the sheet can be reconciled against the ERP.
                if nf in ('quantity', 'unit_price', 'gross_amount', 'discount',
                          'tax', 'net_amount', 'target_amount'):
                    continue
                setattr(rec, nf, parse_num(cell(row, nf)))
            for df in DATE_FIELDS:
                setattr(rec, df, parse_date(cell(row, df)))
            batch.append(rec)
            # A cancelled invoice is not revenue.
            if not cancelled:
                total_rev += net
            lo = od if lo is None or od < lo else lo
            hi = od if hi is None or od > hi else hi

            if len(batch) >= 2000:
                SalesRecord.objects.bulk_create(batch, batch_size=1000)
                batch = []
        if batch:
            SalesRecord.objects.bulk_create(batch, batch_size=1000)
    except Exception as e:
        # Only this sheet's rows. The upload belongs to the file, and another
        # sheet may already have loaded cleanly into it.
        upload.records.filter(id__gt=watermark).delete()
        return Response({'error': f'Failed while reading row {row_no}: {e}'}, status=400)

    count = upload.records.filter(id__gt=watermark).count()
    if count == 0:
        return Response({
            'error': 'No usable rows found — every row was missing a valid date.',
            'detected_columns': sorted(col_map.keys()),
            'unrecognised_columns': unknown,
        }, status=400)

    warnings, notes = [], []
    if line_types:
        notes.append(
            'Line types in the Type column: '
            + ', '.join(f'{k} ({v:,})' for k, v in
                        sorted(line_types.items(), key=lambda x: -x[1])[:8]) + '.')
    if non_item_value:
        notes.append(
            f'Rs {non_item_value:,.0f} of the total sits on lines with no product '
            f'on them - freight, rounding and other charges posted straight to a '
            f'ledger account. They count in the sales total, as they do on the '
            f'invoice, but they cannot appear in a product or SKU breakdown. That '
            f'is why those views add up to slightly less than the headline.')
    if cancelled_rows:
        notes.append(
            f'{cancelled_rows} cancelled row(s) were loaded but are excluded from '
            f'every sales figure. They stay in the data so the file still '
            f'reconciles against the ERP.')
    if return_rows:
        if return_positive and not return_negative:
            warnings.append(
                f'{return_rows} return row(s), every one carrying a POSITIVE amount - '
                f'so they are ADDING to sales instead of reducing them. Say the word '
                f'and we will flip the sign on load.')
        elif return_negative and not return_positive:
            notes.append(
                f'{return_rows} return row(s), all negative - they already reduce '
                f'sales, which is right. Nothing to change.')
        else:
            warnings.append(
                f'{return_rows} return row(s): {return_negative:,} negative and '
                f'{return_positive:,} positive. A mix means the export is not '
                f'consistent about which way a return points.')
    if no_date:
        warnings.append(f'{no_date} row(s) skipped — the date column was empty or '
                        f'unreadable. Those sales are NOT in the dashboard.')
    if bad_value:
        warnings.append(f'{bad_value} row(s) had no sales value (treated as 0). '
                        f'Check the amount column in your export.')
    skipped, unrecognised = partition_unknown(unknown)
    if unrecognised:
        shown = ', '.join(unrecognised[:12]) + (' …' if len(unrecognised) > 12 else '')
        warnings.append(f'{len(unrecognised)} column(s) were not recognised and are '
                        f'ignored: {shown}. Rename them to match the template, or ask '
                        f'for them to be added.')
    missing = [d for d in ('state', 'category', 'channel', 'salesperson')
               if d not in col_map]
    if missing:
        notes.append('No ' + ', '.join(missing) + ' column — those breakdown views '
                        'will be empty. Add the column and re-upload to enable them.')


    return Response({
        'message': f'Imported {count:,} rows.',
        'upload_id': upload.id,
        'rows': count,
        'skipped': no_date,
        'total_revenue': _money(total_rev),
        'period': {'from': lo.isoformat() if lo else None,
                   'to': hi.isoformat() if hi else None},
        'detected_columns': sorted(col_map.keys()),
        # Split so the screen can say "skipped on purpose" and "we don't
        # know what this is" differently. Lumping them together made every
        # upload of a normal ERP dump look like sixteen mapping failures.
        'skipped_columns': skipped,
        'unrecognised_columns': unrecognised,
        'cancelled_rows': cancelled_rows,
        'return_rows': return_rows,
        'line_types': line_types,
        'warnings': warnings,
        'notes': notes,
    })


class SalesOverviewView(APIView):
    """Headline KPIs + comparison against the preceding equal-length window."""

    def get(self, request):
        qs, applied = apply_filters(SalesRecord.objects.all(), request)
        agg = qs.aggregate(
            revenue=Sum('net_amount'), qty=Sum('quantity'), target=Sum('target_amount'),
            orders=Count('invoice_no', distinct=True), lines=Count('id'),
            discount=Sum('discount'),
        )
        revenue = _money(agg['revenue'])
        target = _money(agg['target'])
        lo, hi = _period_bounds(qs)

        # Preceding window of identical length, same dimension filters, so the
        # comparison is like-for-like rather than "this quarter vs all history".
        prev = None
        if lo and hi:
            span = (hi - lo).days + 1
            p_hi = lo - timedelta(days=1)
            p_lo = p_hi - timedelta(days=span - 1)
            pqs, _ = apply_dim_filters(SalesRecord.objects.all(), request)
            prev = (pqs.filter(order_date__gte=p_lo, order_date__lte=p_hi)
                       .aggregate(revenue=Sum('net_amount'), qty=Sum('quantity')))

        prev_rev = _money(prev['revenue']) if prev else 0.0
        customers = qs.exclude(customer_name='').values('customer_name').distinct().count()
        skus = qs.exclude(sku='').values('sku').distinct().count()
        orders = agg['orders'] or 0

        return Response({
            'revenue': revenue,
            'quantity': _money(agg['qty']),
            'orders': orders,
            'lines': agg['lines'] or 0,
            'customers': customers,
            'skus': skus,
            'discount': _money(agg['discount']),
            'avg_order_value': _money(revenue / orders) if orders else 0.0,
            'target': target,
            'achievement_pct': round((revenue / target) * 100, 1) if target else None,
            'gap_to_target': _money(target - revenue) if target else None,
            'prev_revenue': prev_rev,
            'revenue_growth_pct': _pct_change(revenue, prev_rev),
            'quantity_growth_pct': _pct_change(_money(agg['qty']),
                                               _money(prev['qty']) if prev else 0),
            'period': {'from': lo.isoformat() if lo else None,
                       'to': hi.isoformat() if hi else None},
            'filters': applied,
            'has_data': revenue != 0 or (agg['lines'] or 0) > 0,
        })


class SalesBreakdownView(APIView):
    """Group by any whitelisted dimension. `?dim=state&metric=revenue&limit=10`"""

    def get(self, request):
        dim_key = (request.query_params.get('dim') or 'state').strip().lower()
        field = DIMENSIONS.get(dim_key)
        if not field:
            return Response({'error': f'Unknown dimension "{dim_key}".',
                             'available': sorted(DIMENSIONS.keys())}, status=400)
        metric = (request.query_params.get('metric') or 'revenue').strip().lower()
        try:
            limit = max(1, min(200, int(request.query_params.get('limit', 15))))
        except (TypeError, ValueError):
            limit = 15

        qs, applied = apply_filters(SalesRecord.objects.all(), request)
        qs = qs.exclude(**{field: ''})
        rows = (qs.values(field)
                  .annotate(revenue=Sum('net_amount'), quantity=Sum('quantity'),
                            target=Sum('target_amount'), orders=Count('invoice_no', distinct=True),
                            lines=Count('id'))
                  .order_by('-quantity' if metric == 'quantity' else '-revenue'))

        all_rows = list(rows)
        total_rev = sum(float(r['revenue'] or 0) for r in all_rows) or 1.0
        top = all_rows[:limit]
        out = []
        for r in top:
            rev = _money(r['revenue'])
            tgt = _money(r['target'])
            out.append({
                'name': r[field] or '—',
                'revenue': rev,
                'quantity': _money(r['quantity']),
                'orders': r['orders'],
                'lines': r['lines'],
                'target': tgt,
                'achievement_pct': round((rev / tgt) * 100, 1) if tgt else None,
                'share_pct': round((rev / total_rev) * 100, 1),
            })
        others = all_rows[limit:]
        return Response({
            'dimension': dim_key,
            'metric': metric,
            'results': out,
            'total_groups': len(all_rows),
            'others': {
                'count': len(others),
                'revenue': _money(sum(float(r['revenue'] or 0) for r in others)),
            } if others else None,
            'filters': applied,
        })


class SalesTrendView(APIView):
    """Monthly time series, with target and a cumulative running total."""

    def get(self, request):
        qs, applied = apply_filters(SalesRecord.objects.all(), request)
        rows = (qs.values('period')
                  .annotate(revenue=Sum('net_amount'), quantity=Sum('quantity'),
                            target=Sum('target_amount'), orders=Count('invoice_no', distinct=True))
                  .order_by('period'))
        out, running = [], 0.0
        prev_rev = None
        for r in rows:
            rev = _money(r['revenue'])
            running += rev
            tgt = _money(r['target'])
            out.append({
                'period': r['period'].isoformat(),
                'label': r['period'].strftime('%b %Y'),
                'revenue': rev,
                'quantity': _money(r['quantity']),
                'orders': r['orders'],
                'target': tgt,
                'achievement_pct': round((rev / tgt) * 100, 1) if tgt else None,
                'cumulative': round(running, 2),
                'mom_growth_pct': _pct_change(rev, prev_rev) if prev_rev is not None else None,
            })
            prev_rev = rev

        best = max(out, key=lambda r: r['revenue']) if out else None
        worst = min(out, key=lambda r: r['revenue']) if out else None
        return Response({
            'results': out, 'months': len(out),
            'best_month': best, 'worst_month': worst,
            'filters': applied,
        })


class SalesForecastView(APIView):
    """Forecast future monthly revenue from the filtered history."""

    def get(self, request):
        try:
            periods = max(1, min(24, int(request.query_params.get('periods', 6))))
        except (TypeError, ValueError):
            periods = 6
        metric = (request.query_params.get('metric') or 'revenue').strip().lower()
        agg_field = 'quantity' if metric == 'quantity' else 'net_amount'

        qs, applied = apply_filters(SalesRecord.objects.all(), request)
        rows = (qs.values('period').annotate(v=Sum(agg_field)).order_by('period'))
        points = [(r['period'], float(r['v'] or 0)) for r in rows]

        result = forecast_series(points, periods=periods)
        result['metric'] = metric
        result['history'] = [{'period': d.isoformat(), 'label': d.strftime('%b %Y'),
                              'value': round(v, 2)} for d, v in points]
        result['filters'] = applied

        hist_total = sum(v for _, v in points)
        if points and result.get('points'):
            # Compare like with like: the same number of months, most recent first.
            n = min(periods, len(points))
            recent = sum(v for _, v in points[-n:])
            proj = sum(p['value'] for p in result['points'][:n])
            result['vs_recent'] = {
                'months': n,
                'recent_total': round(recent, 2),
                'projected_total': round(proj, 2),
                'change_pct': _pct_change(proj, recent),
            }
        result['history_total'] = round(hist_total, 2)
        return Response(result)


class SalesFiltersView(APIView):
    """Distinct values for every filter, so the UI can populate its dropdowns."""

    def get(self, request):
        # Cancelled rows are excluded from every figure, so offering their
        # values here would put a customer in the dropdown that returns an
        # empty dashboard when picked.
        qs = SalesRecord.objects.exclude(is_cancelled=True)
        out = {}
        for f in FILTERABLE:
            vals = (qs.exclude(**{f: ''}).values_list(f, flat=True)
                      .order_by(f).distinct()[:500])
            out[f] = list(vals)
        lo, hi = _period_bounds(qs)
        out['date_range'] = {'from': lo.isoformat() if lo else None,
                             'to': hi.isoformat() if hi else None}
        out['dimensions'] = sorted(DIMENSIONS.keys())
        return Response(out)


class SalesInsightsView(APIView):
    """Auto-generated written observations — the 'so what' the numbers imply.

    Kept server-side so the same wording appears everywhere the data is shown."""

    def get(self, request):
        qs, applied = apply_filters(SalesRecord.objects.all(), request)
        insights = []

        total = _money(qs.aggregate(v=Sum('net_amount'))['v'])
        if not total:
            return Response({'insights': [], 'filters': applied})

        def top_of(field, label, min_share=0):
            rows = (qs.exclude(**{field: ''}).values(field)
                      .annotate(v=Sum('net_amount')).order_by('-v')[:3])
            rows = [r for r in rows if r['v']]
            if not rows:
                return None
            share = float(rows[0]['v']) / total * 100
            if share < min_share:
                return None
            return rows, share

        # Concentration risk — one region/customer carrying too much of the book.
        for field, noun in (('state', 'state'), ('customer_name', 'customer'),
                            ('category', 'category')):
            got = top_of(field, noun)
            if not got:
                continue
            rows, share = got
            name = rows[0][field]
            if share >= 40:
                insights.append({
                    'type': 'risk',
                    'title': f'Heavy concentration in one {noun}',
                    'body': f'{name} alone accounts for {share:.0f}% of sales. '
                            f'That is a single point of failure — losing it would take a '
                            f'large share of revenue with it.',
                })
            elif share >= 25:
                insights.append({
                    'type': 'info',
                    'title': f'Top {noun}: {name}',
                    'body': f'{name} contributes {share:.0f}% of total sales — the largest '
                            f'single {noun} in this selection.',
                })

        # Target achievement
        tgt = _money(qs.aggregate(v=Sum('target_amount'))['v'])
        if tgt:
            ach = total / tgt * 100
            if ach >= 100:
                insights.append({
                    'type': 'win',
                    'title': f'Target exceeded — {ach:.0f}%',
                    'body': f'Sales of {total:,.0f} against a target of {tgt:,.0f}, '
                            f'ahead by {total - tgt:,.0f}.',
                })
            else:
                insights.append({
                    'type': 'risk' if ach < 80 else 'info',
                    'title': f'Target achievement at {ach:.0f}%',
                    'body': f'Short of target by {tgt - total:,.0f}. '
                            f'{"Well behind plan — worth investigating by region." if ach < 80 else "Within reach of plan."}',
                })
            # Who is dragging
            lag = (qs.exclude(salesperson='').values('salesperson')
                     .annotate(rev=Sum('net_amount'), t=Sum('target_amount'))
                     .filter(t__gt=0).order_by('rev'))
            lag = [r for r in lag if float(r['rev']) / float(r['t']) < 0.7][:3]
            if lag:
                names = ', '.join(f"{r['salesperson']} ({float(r['rev'])/float(r['t'])*100:.0f}%)"
                                  for r in lag)
                insights.append({
                    'type': 'risk',
                    'title': 'Sales people below 70% of target',
                    'body': f'{names}. These are the biggest gaps to close.',
                })

        # Momentum from the monthly series
        months = list(qs.values('period').annotate(v=Sum('net_amount')).order_by('period'))
        if len(months) >= 4:
            recent = [float(m['v'] or 0) for m in months[-3:]]
            earlier = [float(m['v'] or 0) for m in months[-6:-3]] or recent
            r_avg, e_avg = sum(recent) / len(recent), sum(earlier) / len(earlier)
            if e_avg:
                delta = (r_avg - e_avg) / e_avg * 100
                if abs(delta) >= 10:
                    insights.append({
                        'type': 'win' if delta > 0 else 'risk',
                        'title': f'Momentum {"rising" if delta > 0 else "falling"} '
                                 f'{abs(delta):.0f}%',
                        'body': f'The last 3 months average {r_avg:,.0f} vs {e_avg:,.0f} in the '
                                f'3 months before — a clear '
                                f'{"upswing" if delta > 0 else "slowdown"}.',
                    })

        # Discount pressure
        disc = _money(qs.aggregate(v=Sum('discount'))['v'])
        gross = _money(qs.aggregate(v=Sum('gross_amount'))['v'])
        if gross and disc / gross * 100 >= 12:
            insights.append({
                'type': 'risk',
                'title': f'Discounting at {disc / gross * 100:.0f}% of gross',
                'body': f'{disc:,.0f} given away against {gross:,.0f} gross. '
                        f'High discount intensity erodes margin even when top-line looks healthy.',
            })

        # Dormant customers — bought before, nothing recently.
        lo, hi = _period_bounds(qs)
        if hi:
            cutoff = hi - timedelta(days=90)
            recent_c = set(qs.filter(order_date__gt=cutoff)
                             .exclude(customer_name='')
                             .values_list('customer_name', flat=True).distinct())
            all_c = set(qs.exclude(customer_name='')
                          .values_list('customer_name', flat=True).distinct())
            dormant = all_c - recent_c
            if dormant and len(all_c) >= 5:
                insights.append({
                    'type': 'risk',
                    'title': f'{len(dormant)} customer(s) inactive for 90+ days',
                    'body': f'{", ".join(sorted(dormant)[:5])}'
                            f'{" and others" if len(dormant) > 5 else ""} have not ordered in the '
                            f'last 90 days of this period. Worth a win-back call.',
                })

        return Response({'insights': insights, 'filters': applied})


class SalesUploadsView(APIView):
    """List uploads; delete one (rolls back a bad file) or all."""

    def get(self, request):
        ups = SalesUpload.objects.all()[:100]
        return Response({'results': [{
            'id': u.id, 'filename': u.filename, 'rows': u.row_count,
            'notes': u.notes,
            'skipped': u.skipped_rows, 'revenue': _money(u.total_revenue),
            'period_start': u.period_start.isoformat() if u.period_start else None,
            'period_end': u.period_end.isoformat() if u.period_end else None,
            'warnings': u.warnings, 'status': u.status,
            'created_at': u.created_at.isoformat(),
        } for u in ups],
            'total_rows': SalesRecord.objects.count(),
            'count': SalesUpload.objects.count()})

    def delete(self, request):
        up_id = request.query_params.get('id')
        if up_id:
            try:
                u = SalesUpload.objects.get(id=up_id)
            except SalesUpload.DoesNotExist:
                return Response({'error': 'Upload not found'}, status=404)
            n = u.records.count()
            u.delete()   # cascades to its rows
            # Removing the invoice data hands the figures back to the AOP
            # sheet, if one is loaded — otherwise the dashboard would go to
            # zero while a perfectly good plan file sat in the table.
            sync_actual_source()
            return Response({'message': f'Removed upload "{u.filename}" and {n:,} row(s).',
                             'deleted': n})
        n = SalesRecord.objects.count()
        SalesRecord.objects.all().delete()
        SalesUpload.objects.all().delete()
        return Response({'message': f'Cleared all sales data ({n:,} row(s)).', 'deleted': n})


class SalesExportView(APIView):
    """Export the current filtered view as Excel — summary + per-dimension sheets."""

    def get(self, request):
        qs, applied = apply_filters(SalesRecord.objects.all(), request)
        wb = openpyxl.Workbook()
        from openpyxl.styles import Font as F, PatternFill as P

        def sheet(title, header, rows):
            ws = wb.create_sheet(title[:31])
            ws.append(header)
            for c in ws[1]:
                c.font = F(bold=True, color='FFFFFF')
                c.fill = P(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
            for r in rows:
                ws.append(r)
            for i, _ in enumerate(header, 1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 22
            ws.freeze_panes = 'A2'

        agg = qs.aggregate(rev=Sum('net_amount'), qty=Sum('quantity'),
                           tgt=Sum('target_amount'), lines=Count('id'))
        ws = wb.active
        ws.title = 'Summary'
        ws.append(['SalesIQ Export'])
        ws['A1'].font = F(bold=True, size=14)
        ws.append([])
        for k, v in (('Total Revenue', _money(agg['rev'])),
                     ('Total Quantity', _money(agg['qty'])),
                     ('Total Target', _money(agg['tgt'])),
                     ('Rows', agg['lines'] or 0)):
            ws.append([k, v])
        ws.append([])
        ws.append(['Filters applied'])
        for k, v in (applied or {'(none)': ''}).items():
            ws.append([k, ', '.join(v) if isinstance(v, list) else str(v)])
        ws.column_dimensions['A'].width = 24
        ws.column_dimensions['B'].width = 40

        for dim_key in ('state', 'area', 'category', 'product', 'channel', 'salesperson',
                        'customer'):
            field = DIMENSIONS[dim_key]
            rows = (qs.exclude(**{field: ''}).values(field)
                      .annotate(rev=Sum('net_amount'), qty=Sum('quantity'),
                                tgt=Sum('target_amount'))
                      .order_by('-rev')[:500])
            if not rows:
                continue
            sheet(dim_key.title(), [dim_key.title(), 'Revenue', 'Quantity', 'Target',
                                    'Achievement %'],
                  [[r[field], _money(r['rev']), _money(r['qty']), _money(r['tgt']),
                    round(float(r['rev'] or 0) / float(r['tgt']) * 100, 1) if r['tgt'] else '']
                   for r in rows])

        trend = (qs.values('period').annotate(rev=Sum('net_amount'), qty=Sum('quantity'),
                                              tgt=Sum('target_amount')).order_by('period'))
        if trend:
            sheet('Monthly Trend', ['Month', 'Revenue', 'Quantity', 'Target'],
                  [[r['period'].strftime('%b %Y'), _money(r['rev']), _money(r['qty']),
                    _money(r['tgt'])] for r in trend])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        resp = HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp['Content-Disposition'] = 'attachment; filename="SalesIQ_Export.xlsx"'
        return resp

