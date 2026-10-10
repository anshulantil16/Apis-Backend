"""The morning mail: who it is addressed to, and what stops it going.

The failure this file exists for is not a crash. It is a mail that arrives,
correctly formatted, signed off and plausible, carrying somebody else's
numbers under the reader's own name.
"""
import io

from django.core import mail as DJMAIL
from django.test import override_settings

from .models import ReportRecipient, ReviewSnapshot
from .tests import (TestCase, Client, GTR01, GTR04B, REVIEW_HEADERS,
                    review_workbook, upload)
from .views.recipients import candidates, match_head, unique_key


# The handover line, as it reads on the real sheet: ARUN MISHRA stands in for
# ARNAB GHOSH, and his REGION cell is merged upward -- so GTR04 A is the
# region on BOTH rows and names neither of them.
GTR04A = ['GT', 'GTR04 A', 'ARNAB GHOSH', 18, 40.00, 60.00, 3.00, 37.12,
          37.12, 0.62, -0.07, 22.88, 400.00, 250.00, 0.63, 150.00,
          900.00, 250.00, 0.28, 650.00]
HANDOVER = ['GT', 'GTR04 A', 'ARUN MISHRA ( ARNAB GHOSH)', 9, 10.00, 20.00,
            1.00, 12.00, 12.00, 0.60, 0.20, 8.00, 100.00, 60.00, 0.60, 40.00,
            240.00, 60.00, 0.25, 180.00]

MAIL_ON = dict(EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend',
               DEFAULT_FROM_EMAIL='reports@apisindia.com',
               EMAIL_HOST_PASSWORD='not-a-real-password')


def rows():
    return [r for r in ReviewSnapshot.objects.first().rows.all()
            if not r.is_total]


class AKeyHasToNameOnePerson(TestCase):
    """A region is not unique on this sheet."""

    def setUp(self):
        upload(review_workbook([GTR01, GTR04A, HANDOVER]))

    def test_the_shared_region_is_on_two_rows(self):
        self.assertEqual(len(candidates('GTR04 A', rows())), 2)

    def test_a_key_naming_two_people_is_not_a_match(self):
        """It used to resolve to whichever came first, so the handover line's
        recipient would have been sent Arnab Ghosh's numbers."""
        self.assertIsNone(match_head('GTR04 A', rows()))

    def test_the_template_hands_out_a_key_that_names_one_row(self):
        """Not the region, where the region is shared -- the name, which is
        the only thing that tells the two apart."""
        rs = rows()
        keys = {r.head_name.upper(): unique_key(r, rs) for r in rs}
        self.assertEqual(keys['MOHINDER SHARMA'], 'GTR01')
        # The region is shared, so the key falls back to the name -- the only
        # thing that tells the two GTR04 A rows apart.
        self.assertEqual(keys['ARNAB GHOSH'].upper(), 'ARNAB GHOSH')
        self.assertEqual(keys['ARUN MISHRA ( ARNAB GHOSH)'].upper(),
                         'ARUN MISHRA ( ARNAB GHOSH)')
        for name, key in keys.items():
            self.assertEqual([r.head_name.upper() for r in candidates(key, rs)],
                             [name])

    def test_every_key_in_the_workbook_resolves_to_exactly_one_row(self):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(
            self.client.get('/api/sales/recipients/template/').content))
        rs = rows()
        seen = 0
        for row in wb.worksheets[0].iter_rows(min_row=2, values_only=True):
            if row[0] == 'head' and row[1]:
                self.assertEqual(len(candidates(row[1], rs)), 1, row)
                seen += 1
        self.assertEqual(seen, 3)

    def test_the_report_refuses_an_ambiguous_key_rather_than_guessing(self):
        r = self.client.get('/api/sales/review/report/?head=GTR04%20A')
        self.assertEqual(r.status_code, 400)
        self.assertIn('2 rows', r.json()['error'])

    def test_the_list_says_which_two_people_a_key_could_mean(self):
        ReportRecipient.objects.create(role='head', head_key='GTR04 A',
                                       email='x@apisindia.com', name='X')
        rec = self.client.get('/api/sales/recipients/').json()['recipients'][0]
        self.assertFalse(rec['matched'])
        self.assertEqual(sorted(n.upper() for n in rec['ambiguous']),
                         ['ARNAB GHOSH', 'ARUN MISHRA ( ARNAB GHOSH)'])


@override_settings(**MAIL_ON)
class TheMorningSend(TestCase):

    def setUp(self):
        upload(review_workbook([GTR01, GTR04B]))
        ReportRecipient.objects.create(role='head', head_key='GTR01',
                                       name='Mohinder Sharma',
                                       email='mohinder@apisindia.com')
        ReportRecipient.objects.create(role='head', head_key='GTR04 B',
                                       name='Gulshan Kumar',
                                       email='gulshan@apisindia.com')

    def send(self, **body):
        return self.client.post('/api/sales/mail/send/', body,
                                content_type='application/json')

    def test_the_outbox_shows_what_would_go_without_sending_it(self):
        d = self.client.get('/api/sales/mail/outbox/').json()
        self.assertEqual(d['ready'], 2)
        self.assertEqual(len(DJMAIL.outbox), 0)
        self.assertTrue(all(r['attachments'] for r in d['recipients']))

    def test_each_head_is_sent_their_own_report_and_nobody_elses(self):
        self.send()
        self.assertEqual(len(DJMAIL.outbox), 2)
        by_to = {m.to[0]: m for m in DJMAIL.outbox}
        mine = by_to['mohinder@apisindia.com']
        html = mine.alternatives[0][0]
        self.assertIn('GTR01', html)
        self.assertNotIn('GULSHAN', html.upper())
        names = [n for n, _, _ in mine.attachments]
        self.assertEqual(len(names), 1)
        self.assertIn('GTR01', names[0])
        self.assertNotIn('GULSHAN', names[0].upper())

    def test_a_test_run_goes_to_one_address_and_says_so(self):
        self.send(test_to='anshul@apisindia.com')
        self.assertEqual(len(DJMAIL.outbox), 2)
        self.assertTrue(all(m.to == ['anshul@apisindia.com']
                            for m in DJMAIL.outbox))
        self.assertTrue(all(m.subject.startswith('[TEST')
                            for m in DJMAIL.outbox))

    def test_a_test_run_does_not_mark_the_morning_mail_as_sent(self):
        """Otherwise the screen says sixteen people have today's numbers when
        the only inbox touched was your own."""
        self.send(test_to='anshul@apisindia.com')
        self.assertFalse(ReportRecipient.objects.exclude(
            last_sent_at=None).exists())
        self.send()
        self.assertEqual(ReportRecipient.objects.exclude(
            last_sent_at=None).count(), 2)

    def test_one_broken_recipient_stops_the_whole_run(self):
        """Half a run going out is the worst outcome available: nobody can
        say afterwards who holds the morning's numbers."""
        ReportRecipient.objects.create(role='head', head_key='GTR99',
                                       name='Nobody', email='n@apisindia.com')
        r = self.send()
        self.assertEqual(r.status_code, 400)
        self.assertEqual(len(DJMAIL.outbox), 0)
        self.assertEqual(r.json()['failed'][0]['email'], 'n@apisindia.com')

    def test_the_rest_can_be_sent_once_that_is_a_decision(self):
        ReportRecipient.objects.create(role='head', head_key='GTR99',
                                       name='Nobody', email='n@apisindia.com')
        d = self.send(skip_broken=True).json()
        self.assertEqual(d['sent'], 2)
        self.assertEqual(len(DJMAIL.outbox), 2)

    def test_a_manager_gets_one_report_covering_all_of_it(self):
        """Not the roll-up plus one file per head. For a manager holding the
        whole channel that arrived as fourteen attachments -- a strip of
        thumbnails to scroll through rather than a report to read -- and the
        other thirteen were the same figures the roll-up already carried."""
        ReportRecipient.objects.create(role='manager', name='North',
                                       email='mgr@apisindia.com',
                                       regions=['GTR01', 'GTR04 B'])
        self.send(recipients=[ReportRecipient.objects.get(
            email='mgr@apisindia.com').id])
        m = DJMAIL.outbox[0]
        self.assertEqual(len(m.attachments), 1)
        self.assertIn('North', m.attachments[0][0])
        self.assertIn('GTR01', m.alternatives[0][0])
        self.assertIn('GTR04 B', m.alternatives[0][0])

    def test_a_manager_covering_nothing_is_refused_not_sent_an_empty_report(self):
        ReportRecipient.objects.create(role='manager', name='Empty',
                                       email='e@apisindia.com', regions=[])
        r = self.send(recipients=[ReportRecipient.objects.get(
            email='e@apisindia.com').id])
        self.assertEqual(r.status_code, 400)
        self.assertEqual(len(DJMAIL.outbox), 0)

    def test_an_inactive_recipient_is_not_sent_to(self):
        ReportRecipient.objects.filter(head_key='GTR01').update(is_active=False)
        self.send()
        self.assertEqual([m.to[0] for m in DJMAIL.outbox],
                         ['gulshan@apisindia.com'])

    def test_a_reader_cannot_fire_the_send(self):
        """Granting somebody the dashboard is not granting them the ability to
        mail the sales leadership."""
        from .views.auth import issue_session
        c = Client()
        c.defaults['HTTP_X_SALESIQ_SESSION'] = issue_session('reader@apisindia.com')
        r = c.post('/api/sales/mail/send/', {},
                   content_type='application/json')
        self.assertIn(r.status_code, (401, 403))
        self.assertEqual(len(DJMAIL.outbox), 0)


class WithNoSendingAccountConfigured(TestCase):

    def setUp(self):
        upload(review_workbook([GTR01]))
        ReportRecipient.objects.create(role='head', head_key='GTR01',
                                       name='Mohinder Sharma',
                                       email='mohinder@apisindia.com')

    @override_settings(DEFAULT_FROM_EMAIL='', EMAIL_HOST_PASSWORD='')
    def test_it_says_so_rather_than_failing_against_smtp(self):
        r = self.client.post('/api/sales/mail/send/', {},
                             content_type='application/json')
        self.assertEqual(r.status_code, 400)
        self.assertIn('.env', r.json()['error'])
        self.assertFalse(self.client.get(
            '/api/sales/mail/outbox/').json()['configured'])


class AnIndividualReportIsOnlyTheirOwn(TestCase):
    """The report goes to one person.

    A board listing eleven colleagues by name and figure hands every reader
    their peers' numbers. That is somebody else's information, and circulating
    it to sixteen inboxes every morning is not a thing to do by accident.
    """

    def setUp(self):
        upload(review_workbook([GTR01, GTR04A, GTR04B, HANDOVER]))

    def report_for(self, region):
        rs = rows()
        me = next(r for r in rs if r.region == region)
        r = self.client.get('/api/sales/review/report/file/?row=%d' % me.id)
        self.assertEqual(r.status_code, 200)
        return me, rs, r.content.decode('utf-8')

    def test_no_other_head_is_named_anywhere_in_it(self):
        me, rs, html = self.report_for('GTR01')
        for other in rs:
            if other.id == me.id:
                continue
            self.assertNotIn(other.head_name, html,
                             '%s is named in %s report'
                             % (other.head_name, me.head_name))
            # The surname on its own too: "Kumar" in a caption is still a
            # colleague's figure attached to a colleague.
            last = other.head_name.split()[-1].strip('()')
            if len(last) > 3 and last.lower() not in me.head_name.lower():
                self.assertNotIn(last, html)

    def test_it_still_says_where_they_stand(self):
        """Taking the names out must not take the standing out -- that is
        the reader's own position and the point of the section."""
        _, _, html = self.report_for('GTR01')
        flat = ' '.join(html.split())           # the markup wraps its lines
        self.assertIn('How you compare', flat)
        self.assertIn('of 4 this month', flat)
        self.assertIn('middle of GT', flat)

    def test_the_spread_is_still_drawn(self):
        """One mark per territory, so the shape of the channel is visible
        without anybody being named."""
        _, rs, html = self.report_for('GTR01')
        # By the literal colour, because nothing inside an <svg> may use a
        # CSS variable -- WeasyPrint renders those black.
        from .report import C_PEER
        self.assertEqual(html.count('fill="%s"' % C_PEER), len(rs) - 1)

    def test_the_manager_report_does_name_them(self):
        """There the team IS the subject, and a roll-up that will not say
        which territory is behind is useless to the person who has to act."""
        regions = ','.join(sorted({r.region for r in rows()}))
        r = self.client.get('/api/sales/review/team/file/?regions=' + regions)
        self.assertEqual(r.status_code, 200)
        html = r.content.decode('utf-8')
        for row in rows():
            self.assertIn(row.head_name, html)


class TheDateTheSheetDescribes(TestCase):
    """It is not written inside the file -- only the month is, in the column
    headers. A sheet closing on the 30th of September uploaded on the 9th of
    October was stamped the 9th of October, and that date went out in the
    subject line of every mail built from it."""

    def test_it_is_read_off_the_name_the_business_writes_on_it(self):
        upload(review_workbook([GTR01]), name="Primary Master till 30th Sept'2026.xlsx")
        self.assertEqual(ReviewSnapshot.objects.first().as_of_date,
                         __import__('datetime').date(2026, 9, 30))

    def test_the_forms_the_business_actually_uses(self):
        from .review import date_from_name as f
        from datetime import date
        self.assertEqual(f("Primary Master till 30th Sept'2026.xlsx"), date(2026, 9, 30))
        self.assertEqual(f('Primary Master till 30.09.2026.xlsx'), date(2026, 9, 30))
        self.assertEqual(f('Sales 30-Sep-26.xlsx'), date(2026, 9, 30))
        self.assertEqual(f('Primary Master till 5 April 2026.xlsx'), date(2026, 4, 5))
        # ISO, which the day-first pattern would read as the 26th of 2030.
        self.assertEqual(f('report_2026-09-30.xlsx'), date(2026, 9, 30))

    def test_something_that_is_not_a_date_is_not_guessed_at(self):
        """Falling back to the upload day is better than inventing a day."""
        from .review import date_from_name as f
        self.assertIsNone(f('Primary Master till 31st Feb 2026.xlsx'))
        self.assertIsNone(f('Region Summary v2.xlsx'))
        self.assertIsNone(f(''))

    def test_an_explicit_as_of_still_wins(self):
        upload(review_workbook([GTR01]),
               name="Primary Master till 30th Sept'2026.xlsx", as_of='2026-10-01')
        self.assertEqual(ReviewSnapshot.objects.first().as_of_date,
                         __import__('datetime').date(2026, 10, 1))

    def test_the_mail_carries_that_date_and_not_the_upload_day(self):
        upload(review_workbook([GTR01]), name="Primary Master till 30th Sept'2026.xlsx")
        ReportRecipient.objects.create(role='head', head_key='GTR01',
                                       name='Mohinder Sharma',
                                       email='m@apisindia.com')
        d = self.client.get('/api/sales/mail/outbox/').json()
        self.assertIn('30.09.26', d['recipients'][0]['subject'])


@override_settings(**MAIL_ON)
class TheMailBody(TestCase):

    def setUp(self):
        upload(review_workbook([GTR01]), name="Primary Master till 30th Sept'2026.xlsx")
        ReportRecipient.objects.create(role='head', head_key='GTR01',
                                       name='Mohinder Sharma',
                                       email='m@apisindia.com')
        self.client.post('/api/sales/mail/send/', {},
                         content_type='application/json')
        self.html = DJMAIL.outbox[0].alternatives[0][0]

    def test_it_reads_like_the_mail_the_business_already_sends(self):
        self.assertIn('Subzone-wise and ASM/TSM-wise B2C Primary Sales Report',
                      self.html)
        self.assertIn("SEPT&#8217;26", self.html)
        self.assertIn('(as of 30.09.26)', self.html)

    def test_the_ot_caveat_is_highlighted_the_way_it_is_circulated(self):
        """Yellow is how somebody scanning on a phone finds it. Reformatting
        it into our own house style makes the mail harder to read for
        everybody who already reads it."""
        self.assertIn('background:#FFFF00', self.html)
        self.assertIn('In OT channel Primary = Secondary.', self.html)

    def test_the_cn_note_is_red(self):
        self.assertIn('color:#FF0000', self.html)
        self.assertIn('Sales Return-Damage Expiry &amp; Good SR', self.html)

    def test_the_attachment_is_not_sent_as_its_own_source_code(self):
        """An .html attachment shows in Gmail as a wall of CSS, and only a
        download gets the reader the report."""
        from . import pdf as PDF
        name, _, mime = (DJMAIL.outbox[0].attachments[0][0],
                         None, DJMAIL.outbox[0].attachments[0][2])
        if PDF.available():
            self.assertTrue(name.endswith('.pdf'), name)
            self.assertEqual(mime, 'application/pdf')
        else:
            # Honest about the fallback rather than silently shipping the
            # thing this test exists to prevent.
            self.assertTrue(name.endswith('.html'), name)
            self.assertEqual(
                self.client.get('/api/sales/mail/outbox/').json()
                ['attachment_format'], 'html')


class WhenTheNameAndTheSheetDisagree(TestCase):
    """The month is the sheet's own statement -- its headers read "MTD Sep-26
    PRI SALES". A September workbook saved under an April name is a September
    workbook with the wrong name on it, and dating its figures to April is
    exactly the kind of quiet relabelling that makes a whole report wrong."""

    def test_the_name_does_not_move_the_month(self):
        upload(review_workbook([GTR01]), name='Primary Master till 01.04.2026.xlsx')
        snap = ReviewSnapshot.objects.first()
        self.assertEqual(snap.as_of_month.month, 9)      # the columns say Sep

    def test_a_contradicting_date_is_refused_not_taken(self):
        """It used to be taken silently, so the mail read "SEPT'26 (as of
        01.04.26)" and left the reader to notice."""
        from datetime import date
        upload(review_workbook([GTR01]), name='Primary Master till 01.04.2026.xlsx')
        snap = ReviewSnapshot.objects.first()
        self.assertNotEqual(snap.as_of_date, date(2026, 4, 1))

    def test_and_it_is_said_out_loud(self):
        r = upload(review_workbook([GTR01]),
                   name='Primary Master till 01.04.2026.xlsx')
        said = ' '.join(r.json().get('warnings') or [])
        self.assertIn('01 April 2026', said)
        self.assertIn('September', said)

    def test_a_name_that_agrees_is_used_as_before(self):
        from datetime import date
        upload(review_workbook([GTR01]),
               name="Primary Master till 30th Sept'2026.xlsx")
        snap = ReviewSnapshot.objects.first()
        self.assertEqual(snap.as_of_date, date(2026, 9, 30))
        self.assertEqual(snap.warnings, [])

    def test_april_s_own_sheet_reads_april_without_being_told(self):
        """Which is the case that actually happens: when April's file is
        uploaded its columns say Apr-26, and nothing has to be renamed."""
        apr = [h.replace("Sep-26", "Apr-26").replace("SEPT'26", "APR'26")
               if isinstance(h, str) else h for h in REVIEW_HEADERS]
        upload(review_workbook([GTR01], headers=apr),
               name='Primary Master till 30.04.2026.xlsx')
        snap = ReviewSnapshot.objects.first()
        self.assertEqual((snap.as_of_month.year, snap.as_of_month.month),
                         (2026, 4))
        from datetime import date
        self.assertEqual(snap.as_of_date, date(2026, 4, 30))
        self.assertEqual(snap.warnings, [])

    def test_the_mail_then_says_april(self):
        apr = [h.replace("Sep-26", "Apr-26").replace("SEPT'26", "APR'26")
               if isinstance(h, str) else h for h in REVIEW_HEADERS]
        upload(review_workbook([GTR01], headers=apr),
               name='Primary Master till 30.04.2026.xlsx')
        ReportRecipient.objects.create(role='head', head_key='GTR01',
                                       name='Mohinder Sharma',
                                       email='m@apisindia.com')
        from . import mail as MAIL
        from .views.review import snapshot_json
        snap = snapshot_json(ReviewSnapshot.objects.first())
        self.assertEqual(MAIL.month_label(snap), 'APR&#8217;26')
        self.assertEqual(MAIL.as_of(snap), '30.04.26')


class AManagerCoveringTheWholeChannel(TestCase):
    """Thirteen regions typed into one cell, which is what somebody filling
    in a spreadsheet actually does."""

    def setUp(self):
        upload(review_workbook([GTR01, GTR04A, GTR04B]))

    def send_list(self, covered):
        rows = [['role', 'key', 'region', 'name', 'email', 'regions_covered'],
                ['manager', 'HO0239', 'GT', 'Anshul Antil',
                 'anshul@apisindia.com', covered]]
        import openpyxl
        wb = openpyxl.Workbook()
        for r in rows:
            wb.active.append(r)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        buf.name = 'recipients.xlsx'
        return self.client.post('/api/sales/recipients/import/',
                                {'file': buf}).json()

    def test_alt_enter_inside_one_cell_is_a_list(self):
        """It used to come back as a single region name, no such region
        existed, and the manager was silently left off."""
        d = self.send_list('GTR01\nGTR04 A\nGTR04 B')
        self.assertEqual(d['skipped'], [])
        self.assertEqual(d['added'], 1)
        rec = ReportRecipient.objects.get(role='manager')
        self.assertEqual(sorted(rec.regions), ['GTR01', 'GTR04 A', 'GTR04 B'])

    def test_semicolons_and_commas_too(self):
        for sep in (';', ',', '; ', '\r\n'):
            ReportRecipient.objects.all().delete()
            d = self.send_list(sep.join(['GTR01', 'GTR04 B']))
            self.assertEqual(d['skipped'], [], sep)
            self.assertEqual(ReportRecipient.objects.get(role='manager').regions,
                             ['GTR01', 'GTR04 B'])

    def test_a_region_listed_twice_is_carried_once(self):
        """GTR04 A is on two rows of the sheet, so it is natural to write it
        twice -- and a group total that counted it twice would be wrong."""
        self.send_list('GTR01\nGTR04 A\nGTR04 A\nGTR04 B')
        self.assertEqual(ReportRecipient.objects.get(role='manager').regions,
                         ['GTR01', 'GTR04 A', 'GTR04 B'])

    def test_the_sheet_s_own_spelling_is_what_gets_stored(self):
        self.send_list('gtr01;gtr04 b')
        self.assertEqual(ReportRecipient.objects.get(role='manager').regions,
                         ['GTR01', 'GTR04 B'])

    def test_a_region_that_is_not_on_the_sheet_says_what_is(self):
        d = self.send_list('GTR01\nGTR77')
        self.assertEqual(d['added'], 0)
        self.assertIn('"GTR77"', d['skipped'][0])
        self.assertIn('GTR01', d['skipped'][0])     # what the sheet carries

    def test_the_word_ALL_still_means_every_territory(self):
        self.send_list('ALL')
        rec = ReportRecipient.objects.get(role='manager')
        self.assertTrue(rec.covers_all)
        self.assertEqual(rec.regions, [])


@override_settings(**MAIL_ON)
class TheManagersCumulatedMail(TestCase):

    def setUp(self):
        upload(review_workbook([GTR01, GTR04A, GTR04B]))
        ReportRecipient.objects.create(
            role='manager', name='Anshul Antil', email='anshul@apisindia.com',
            regions=['GTR01', 'GTR04 A', 'GTR04 B'])
        self.client.post('/api/sales/mail/send/', {},
                         content_type='application/json')
        self.html = DJMAIL.outbox[0].alternatives[0][0]

    def test_the_body_carries_every_territory_he_covers(self):
        for r in ('GTR01', 'GTR04 A', 'GTR04 B'):
            self.assertIn(r, self.html)

    def test_and_one_cumulated_line_under_them(self):
        """Added from the rows he covers, never read off the sheet's own GT
        Total: a manager covering part of a channel has no subtotal there,
        and taking the one that is there would hand him the whole channel."""
        self.assertIn('Group', self.html)
        self.assertIn('3 territories', self.html)

    def test_the_cumulated_figures_are_the_sum_of_his_rows(self):
        from .report import lakh
        total = sum(float(r.mtd_primary) for r in rows()
                    if r.region in ('GTR01', 'GTR04 A', 'GTR04 B'))
        self.assertIn(lakh(total), self.html)

    def test_one_attachment_not_one_per_territory(self):
        self.assertEqual(len(DJMAIL.outbox[0].attachments), 1)

    def test_and_that_one_file_carries_every_territory(self):
        """Which is what makes a single attachment the right answer rather
        than a thing left out."""
        body = DJMAIL.outbox[0].attachments[0][1]
        if isinstance(body, bytes):          # a PDF, where one can be built
            self.skipTest('rendered to PDF; the HTML is checked above')
        for r in ('GTR01', 'GTR04 A', 'GTR04 B'):
            self.assertIn(r, body)


class TwoPeopleInOneGTR(TestCase):
    """A handover line carries the region of the row above it, merged down,
    so the region names two people. The row carries the name as well, and
    that is what tells them apart -- the importer should use it rather than
    send somebody back to download a better list."""

    def setUp(self):
        upload(review_workbook([GTR01, GTR04A, HANDOVER]))
        # The real shape: one of the pair carries an APIS ID and the other
        # does not, which is why the region ends up doing the work.
        for r in rows():
            if r.head_name.upper() == 'ARNAB GHOSH':
                r.head_code = 'SL00979'
                r.save()

    def send(self, lines):
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.append(['role', 'key', 'region', 'name', 'email',
                          'regions_covered'])
        for line in lines:
            wb.active.append(line)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        buf.name = 'recipients.xlsx'
        return self.client.post('/api/sales/recipients/import/',
                                {'file': buf}).json()

    def test_the_shared_region_plus_the_name_is_enough(self):
        d = self.send([
            ['head', 'SL00979', 'GTR04 A', 'Arnab Ghosh', 'a@apisindia.com', ''],
            ['head', 'GTR04 A', 'GTR04 A', 'Arun Mishra ( Arnab Ghosh)',
             'h@apisindia.com', ''],
        ])
        self.assertEqual(d['skipped'], [])
        self.assertEqual(d['added'], 2)

    def test_and_they_are_matched_to_different_rows(self):
        """The whole point. Both resolving to Arnab means the handover line
        is sent his numbers under his own name."""
        self.send([
            ['head', 'SL00979', 'GTR04 A', 'Arnab Ghosh', 'a@apisindia.com', ''],
            ['head', 'GTR04 A', 'GTR04 A', 'Arun Mishra ( Arnab Ghosh)',
             'h@apisindia.com', ''],
        ])
        got = {r['email']: r['matched_to']
               for r in self.client.get('/api/sales/recipients/').json()['recipients']}
        self.assertEqual(got['a@apisindia.com'], 'GTR04 A — Arnab Ghosh')
        self.assertEqual(got['h@apisindia.com'],
                         'GTR04 A — Arun Mishra ( Arnab Ghosh)')

    def test_nobody_is_left_without_a_report(self):
        self.send([
            ['head', 'SL00979', 'GTR04 A', 'Arnab Ghosh', 'a@apisindia.com', ''],
            ['head', 'GTR04 A', 'GTR04 A', 'Arun Mishra ( Arnab Ghosh)',
             'h@apisindia.com', ''],
            ['head', 'GTR01', 'GTR01', 'Mohinder Sharma', 'm@apisindia.com', ''],
        ])
        d = self.client.get('/api/sales/recipients/').json()
        self.assertEqual(d['heads_without_a_recipient'], [])

    def test_it_is_stored_under_a_key_that_can_only_mean_one_row(self):
        """Settled once, at import, rather than every morning when the
        report is built."""
        self.send([['head', 'GTR04 A', 'GTR04 A', 'Arun Mishra ( Arnab Ghosh)',
                    'h@apisindia.com', '']])
        rec = ReportRecipient.objects.get(email='h@apisindia.com')
        self.assertNotEqual(rec.head_key, 'GTR04 A')
        self.assertEqual(len(candidates(rec.head_key, rows())), 1)

    def test_the_old_ambiguous_row_does_not_linger(self):
        """Left behind it would sit on the list unmatched for good."""
        ReportRecipient.objects.create(role='head', head_key='GTR04 A',
                                       email='h@apisindia.com', name='Arun')
        self.send([['head', 'GTR04 A', 'GTR04 A', 'Arun Mishra ( Arnab Ghosh)',
                    'h@apisindia.com', '']])
        self.assertEqual(
            list(ReportRecipient.objects.filter(email='h@apisindia.com')
                 .values_list('head_key', flat=True)),
            ['Arun Mishra ( Arnab Ghosh)'])

    def test_a_name_that_matches_neither_is_reported_not_guessed(self):
        d = self.send([['head', 'GTR04 A', 'GTR04 A', 'Somebody Else',
                        'x@apisindia.com', '']])
        self.assertEqual(d['added'], 0)
        self.assertIn('does not say which', d['skipped'][0])

    def test_the_name_alone_works_too(self):
        """Which is what the template hands out now."""
        d = self.send([['head', 'Arun Mishra ( Arnab Ghosh)', 'GTR04 A',
                        'Arun Mishra ( Arnab Ghosh)', 'h@apisindia.com', '']])
        self.assertEqual(d['skipped'], [])
        self.assertEqual(d['added'], 1)


class WhatTheScreenFlagsAsAProblem(TestCase):
    """A head with no APIS ID used to be flagged in amber on every card. That
    mattered when the ID was the only key. It is now addressed by region, or
    by name where the region is shared, and nothing is held up -- so the tag
    marked working rows as faulty, which is how people learn to ignore a
    colour that is sometimes real."""

    def rows_json(self):
        return self.client.get('/api/sales/review/').json()['rows']

    def test_a_head_with_no_apis_id_is_not_flagged(self):
        upload(review_workbook([GTR01, GTR04A, HANDOVER]),
               name="Primary Master till 30th Sept'2026.xlsx")
        for r in self.rows_json():
            self.assertTrue(r['addressable'], r['head_name'])
            self.assertEqual(r['clashes_with'], [])

    def test_the_handover_line_is_addressed_by_name(self):
        upload(review_workbook([GTR01, GTR04A, HANDOVER]))
        by = {r['head_name'].upper(): r for r in self.rows_json()}
        self.assertEqual(by['MOHINDER SHARMA']['addressed_by'], 'region')
        self.assertEqual(by['ARUN MISHRA ( ARNAB GHOSH)']['addressed_by'], 'name')

    def test_the_upload_no_longer_claims_it_cannot_be_addressed(self):
        r = upload(review_workbook([GTR01, GTR04A, HANDOVER]))
        said = ' '.join((r.json().get('notes') or [])
                        + (r.json().get('warnings') or []))
        self.assertNotIn('cannot be', said)
        self.assertNotIn('APIS ID', said)

    def test_two_rows_that_really_cannot_be_told_apart_are_flagged(self):
        """Same name, same region, no ID on either. Then there genuinely is
        no way to say whose report is whose."""
        twin = list(HANDOVER)
        r = upload(review_workbook([GTR01, HANDOVER, twin]))
        said = ' '.join(r.json().get('warnings') or [])
        self.assertIn('cannot be told apart', said)
        flagged = [x for x in self.rows_json() if not x['addressable']]
        self.assertEqual(len(flagged), 2)
        self.assertTrue(all(x['clashes_with'] for x in flagged))


class TheManagersReportIsWorthOpening(TestCase):
    """It was a bar chart, three tiles and a table -- correct, and thin
    enough that the one question a manager has each week ("where do I spend
    Tuesday?") was left to be worked out from the table."""

    def setUp(self):
        upload(review_workbook([GTR01, GTR04A, GTR04B, HANDOVER]),
               name="Primary Master till 30th Sept'2026.xlsx")

    def report(self):
        regions = ','.join(sorted({r.region for r in rows()}))
        r = self.client.get('/api/sales/review/team/file/?regions='
                            + regions + '&name=Anshul+Antil')
        self.assertEqual(r.status_code, 200)
        return r.content.decode('utf-8')

    def test_it_answers_three_different_questions_not_one(self):
        html = self.report()
        for h in ('How each territory is doing',    # who is behind
                  'Where the gap is',                # where the money is
                  'Which way each is moving'):       # which way they are going
            self.assertIn(h, html)

    def test_the_shortfall_is_shown_in_money_not_only_percent(self):
        """An achievement board read on its own sends a manager to the
        smallest territory they have."""
        html = self.report()
        self.assertIn('lakhs, biggest first', html)
        self.assertIn('to bill', html)

    def test_the_group_gets_the_same_three_gauges_a_head_does(self):
        """One visual language across both documents: a manager who reads
        fourteen of these should not have to learn two."""
        self.assertEqual(self.report().count('A 78 78 0 0 1 178 112'), 6)

    def test_every_territory_is_still_listed_with_a_group_line(self):
        html = self.report()
        for r in rows():
            self.assertIn(r.head_name, html)
        self.assertIn('Group', html)

    def test_the_terms_are_explained_as_they_are_on_a_head_report(self):
        self.assertIn('What these words mean', self.report())


class BothReportsPrintInColour(TestCase):
    """Printed, the report was a white sheet with a blue chip on it --
    correct and characterless, and the first thing anybody sees of it."""

    def setUp(self):
        upload(review_workbook([GTR01, GTR04B]),
               name="Primary Master till 30th Sept'2026.xlsx")

    def both(self):
        me = rows()[0]
        head = self.client.get('/api/sales/review/report/file/?row=%d' % me.id)
        regions = ','.join(sorted({r.region for r in rows()}))
        team = self.client.get('/api/sales/review/team/file/?regions=' + regions)
        return head.content.decode('utf-8'), team.content.decode('utf-8')

    def test_the_masthead_carries_colour_on_both(self):
        for html in self.both():
            self.assertIn('linear-gradient(118deg', html)

    def test_and_the_colour_survives_the_print_path(self):
        """Without this the band comes out white again, which is exactly the
        version this was meant to stop being."""
        for html in self.both():
            self.assertIn('print-color-adjust:exact', html)


class NothingIsDrawnOutsideItsChart(TestCase):
    """A mark outside the viewBox is simply not drawn, and a chart that
    quietly loses its worst bar is worse than one that looks wrong."""

    def setUp(self):
        upload(review_workbook([GTR01, GTR04A, GTR04B, HANDOVER]),
               name="Primary Master till 30th Sept'2026.xlsx")

    def test_every_mark_stays_inside_every_svg(self):
        import re
        regions = ','.join(sorted({r.region for r in rows()}))
        me = rows()[0]
        pages = [
            self.client.get('/api/sales/review/team/file/?regions=' + regions)
                .content.decode('utf-8'),
            self.client.get('/api/sales/review/report/file/?row=%d' % me.id)
                .content.decode('utf-8'),
        ]
        problems = []
        for html in pages:
            for sv in re.finditer(r'<svg viewBox="0 0 ([\d.]+) ([\d.]+)"(.*?)</svg>',
                                  html, re.S):
                W, H, body = float(sv.group(1)), float(sv.group(2)), sv.group(3)
                for m in re.finditer(
                        r'<rect[^>]*?x="([-\d.]+)"[^>]*?width="([-\d.]+)"', body):
                    x, w = float(m.group(1)), float(m.group(2))
                    if w < 0 or x + w > W + 1:
                        problems.append(('rect', x, w, W))
                for m in re.finditer(r'<circle cx="([-\d.]+)" cy="([-\d.]+)"', body):
                    cx, cy = float(m.group(1)), float(m.group(2))
                    if not (0 <= cx <= W and 0 <= cy <= H):
                        problems.append(('circle', cx, cy, W, H))
        self.assertEqual(problems, [])


class AMonthWithMoreReturnedThanSold(TestCase):
    """The sheet writes it as a minus, and it is a real thing. A negative
    width is not drawn by any browser, so the bar vanishes -- and the
    territory in the worst trouble is the one that disappears."""

    def test_the_shortfall_chart_still_draws_every_territory(self):
        upload(review_workbook([GTR01, GTR04B]),
               name="Primary Master till 30th Sept'2026.xlsx")
        # In rupees, as the rows are stored: a figure that rounds to a
        # whisker below zero is how negative zero gets into a width.
        bad = rows()[0]
        bad.mtd_primary = -12.40 * 100000
        bad.save()
        tiny = rows()[1]
        tiny.mtd_primary = -0.4
        tiny.save()
        regions = ','.join(sorted({r.region for r in rows()}))
        html = self.client.get(
            '/api/sales/review/team/file/?regions=' + regions
        ).content.decode('utf-8')
        import re
        self.assertEqual(
            re.findall(r'(?:width|height|r)="(-[\d.]+|nan|NaN|inf)"', html), [])
        # And the figure itself is still printed, whatever its sign.
        self.assertIn('-12.40', html)
        # width="-0.0" is not a valid SVG length either, and max(-0.0, 0) is
        # -0.0 in Python -- which is exactly what a near-zero figure gives.
        self.assertNotIn('"-0.0"', html)


class WhatThePdfActuallyLooksLike(TestCase):
    """Three faults that are invisible on screen and ruin the printed file,
    which is the only version anybody receives."""

    def setUp(self):
        upload(review_workbook([GTR01, GTR04A, GTR04B, HANDOVER]),
               name="Primary Master till 30th Sept'2026.xlsx")

    def pages(self):
        regions = ','.join(sorted({r.region for r in rows()}))
        me = rows()[0]
        return [
            self.client.get('/api/sales/review/team/file/?regions=' + regions)
                .content.decode('utf-8'),
            self.client.get('/api/sales/review/report/file/?row=%d' % me.id)
                .content.decode('utf-8'),
        ]

    def test_no_chart_colour_is_a_css_variable(self):
        """WeasyPrint does not resolve custom properties inside SVG, so a
        fill naming one falls back to black. On screen the charts were green
        and orange; in the PDF that goes out, every bar was black."""
        import re
        for html in self.pages():
            for sv in re.finditer(r'<svg.*?</svg>', html, re.S):
                self.assertNotIn('var(--', sv.group(0))

    def test_every_label_fits_inside_its_own_drawing(self):
        """There is no scrollbar on paper: a label past the viewBox is cut
        off, which is what happened to the handover line's name and to
        "Average of the first 5 months"."""
        import re
        over = []
        for html in self.pages():
            for sv in re.finditer(
                    r'<svg viewBox="0 0 ([\d.]+) ([\d.]+)"(.*?)</svg>', html, re.S):
                W, body = float(sv.group(1)), sv.group(3)
                for t in re.finditer(
                        r'<text x="([-\d.]+)" y="[-\d.]+"([^>]*)>(.*?)</text>',
                        body, re.S):
                    x, attrs = float(t.group(1)), t.group(2)
                    txt = re.sub(r'<[^>]+>', '', t.group(3))
                    txt = txt.replace('&#183;', '.').replace('&mdash;', '-').strip()
                    m = re.search(r'font-size:([\d.]+)px', attrs)
                    # Generously wide on purpose: the PDF falls back to
                    # different fonts from the browser, so a label that just
                    # fits on screen can still run off the page.
                    w = len(txt) * (float(m.group(1)) if m else 11.5) * 0.62
                    if 'text-anchor="end"' in attrs:
                        x0, x1 = x - w, x
                    elif 'text-anchor="middle"' in attrs:
                        x0, x1 = x - w / 2, x + w / 2
                    else:
                        x0, x1 = x, x + w
                    if x0 < -0.5 or x1 > W + 0.5:
                        over.append((round(x0), round(x1), W, txt[:40]))
        self.assertEqual(over, [])

    def test_no_label_is_printed_on_top_of_another_or_on_a_dot(self):
        """A label sitting on the mark it names hides the very thing it is
        there to point at."""
        import re
        html = self.pages()[0]
        m = re.search(r'aria-label="Achievement against growth.*?</svg>', html, re.S)
        if not m:
            self.skipTest('too few territories for the quadrant')
        q = m.group(0)
        boxes = [(float(a) - 8, float(b) - 8, float(a) + 8, float(b) + 8)
                 for a, b in re.findall(r'<circle cx="([-\d.]+)" cy="([-\d.]+)"', q)]
        labels = re.findall(
            r'<text x="([-\d.]+)" y="([-\d.]+)"[^>]*font-size:10px[^>]*>([^<]+)</text>',
            q)
        self.assertEqual(len(labels), len(boxes))   # every territory named
        for x, y, t in labels:
            x, y = float(x), float(y)
            w = len(t) * 5.6 + 4
            box = (x - w / 2, y - 9, x + w / 2, y + 3)
            for o in boxes:
                self.assertTrue(
                    box[2] < o[0] or box[0] > o[2] or box[3] < o[1] or box[1] > o[3],
                    '%s overlaps a mark' % t)
            boxes.append(box)
